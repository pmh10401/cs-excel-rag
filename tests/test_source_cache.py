"""변환 전후의 원문·형식·출처 일치, 손상 파일과 실패 시 원본 보존을 검사합니다."""
import gzip
import json
import tempfile
from datetime import datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from ragkit.records import read_excel, row_revision
from ragkit.source_cache import convert_excel


def main(root):
    with tempfile.TemporaryDirectory() as folder:
        folder = Path(folder)
        source, cache = folder / '원문.xlsx', folder / 'input.rag.jsonl.gz'
        book = Workbook()
        sheet = book.active
        sheet.title = 'CSList'
        columns = ['CS 접수번호', 'line', '현상', '날짜', '시각', '기간', '수량', '실수', '확인', '빈칸']
        sheet.append(['CS 접수번호', 'Line', *columns[2:]])
        sheet.append(['case-1', 'L1', '  순환량 저하\n교환 안 함.  ', datetime(2026, 1, 2, 13, 30), time(8, 40), timedelta(hours=30), 0, 1.25, False, None])
        sheet.append([None] * len(columns))
        sheet.append(['case-2', 'L2', '정상 동작'])
        book.save(source)
        book.close()
        options = dict(required_columns=columns, id_column='CS 접수번호')
        expected = read_excel(source, 'CSList', **options)
        events = []
        result = convert_excel(source, cache, progress=events.append, **options)
        assert result['rows'] == 2
        assert events and events[0]['phase'] == 'opening'
        assert events[-1]['phase'] == 'done' and events[-1]['percent'] == 100
        assert events[-1]['completed'] == 2 and events[-1]['total_rows'] == 3  # 빈 시트 행도 포함
        assert all(e['percent'] != 100 for e in events[:-1]), '파일 저장 전에 완료로 표시했습니다.'
        assert read_excel(cache, 'CSList', **options) == expected
        before = cache.read_bytes()
        with patch.object(Path, 'replace', side_effect=PermissionError('저장 실패')):
            events.clear()
            try:
                convert_excel(source, cache, progress=events.append, **options)
            except PermissionError:
                pass
            else:
                raise AssertionError('저장 실패를 숨겼습니다.')
        assert cache.read_bytes() == before
        assert events[-1]['phase'] == 'saving' and all(e['percent'] != 100 for e in events)
        # 일부 내보내기 도구의 잘못된 사용 범위도 실제 마지막 행을 자르면 안 됩니다.
        original = source.read_bytes()
        with ZipFile(source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries['xl/worksheets/sheet1.xml'] = entries['xl/worksheets/sheet1.xml'].replace(b'A1:J4', b'A1:J2')
        with ZipFile(source, 'w', ZIP_DEFLATED) as archive:
            for name, content in entries.items():
                archive.writestr(name, content)
        events.clear()
        convert_excel(source, cache, progress=events.append, **options)
        assert events[-1]['total_rows'] == 3, '잘못된 Excel 사용 범위를 진행률 분모로 썼습니다.'
        assert read_excel(cache, 'CSList', **options) == expected
        source.write_bytes(original)
        raw = cache.read_bytes()
        source.rename(folder / 'moved.xlsx')
        moved = folder / 'moved.rag.jsonl.gz'
        cache.rename(moved)
        with patch('ragkit.records.load_workbook', side_effect=AssertionError('Excel을 다시 읽었습니다.')):
            actual = read_excel(moved, 'CSList', reader='xlwings', **options)
        assert actual == expected and [row_revision(r) for r in actual] == [row_revision(r) for r in expected]
        assert actual[1]['_excel_location'].endswith('!A4:J4')
        items = [json.loads(line) for line in gzip.decompress(raw).decode().splitlines()]
        corruptions = [raw[:-8], gzip.compress(b'{}\n'), gzip.compress(('\n'.join(json.dumps(x) for x in items[:-1])+'\n').encode())]
        duplicate = json.loads(json.dumps(items))
        duplicate[2]['values'][0] = duplicate[1]['values'][0]
        bad_type = json.loads(json.dumps(items))
        bad_type[1]['typed']['수량']['value'] = 'incorrect'
        for bad in (duplicate, bad_type):
            corruptions.append(gzip.compress(('\n'.join(json.dumps(x) for x in bad)+'\n').encode()))
        for content in corruptions:
            cache.write_bytes(content)
            try:
                read_excel(cache, 'CSList', **options)
            except ValueError:
                pass
            else:
                raise AssertionError('손상된 변환 파일을 허용했습니다.')
        for wrong in ({'required_columns': columns[:-1], 'id_column': 'CS 접수번호'}, {'id_column': 'id'}):
            try:
                read_excel(moved, 'CSList', **wrong)
            except ValueError:
                pass
            else:
                raise AssertionError('양식이 다른 변환 파일을 허용했습니다.')
        source = folder / 'moved.xlsx'
        before = moved.read_bytes()
        book = load_workbook(source)
        book['CSList']['C2'] = '=1+1'
        book.save(source)
        book.close()
        for merged in (False, True):
            if merged:
                book = load_workbook(source)
                book['CSList']['C2'] = '메모'
                book['CSList'].merge_cells('C2:C3')
                book.save(source)
                book.close()
            try:
                events.clear()
                convert_excel(source, moved, progress=events.append, **options)
            except ValueError:
                pass
            else:
                raise AssertionError('수식 또는 병합 셀을 허용했습니다.')
            assert moved.read_bytes() == before
            assert all(e['percent'] != 100 for e in events), '실패한 변환을 완료로 표시했습니다.'
        assert not list(folder.glob('*.tmp'))
    print('PASS: 입력 변환·재사용·출처/자료형·손상 검사·실패 시 기존 파일 보존')


if __name__ == '__main__':
    main(Path(__file__).resolve().parents[1] / 'windows-demo')
