"""28열 다국어 입력이 기존 DB·검토·검색·인용과 함께 동작하는지 검사합니다."""
import copy
import gzip
import json
import sqlite3
import sys
from contextlib import closing
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace as NS
from unittest.mock import Mock, patch

from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel
from ragkit import cli, cslist, english, llm, records, search, semantic, source_cache, storage, web
from ragkit.input_formats import (HISTORY_COLUMNS, HISTORY_SHEET, HISTORY_REWRITE, HISTORY_DETAIL_COLUMNS,
                                 LEGACY_HISTORY_COLUMNS, LEGACY_HISTORY_SHEET, row_schema)


def check_details(root):
    """같은 CS의 부품·내용 차이, 완전 중복, 순서 변경과 구형 DB 재사용을 검사합니다."""
    schema, rules = cslist.load_schema(root / 'cslist_schema.json')
    options = {'required_columns': schema['columns'], 'id_column': schema['id_column']}
    with TemporaryDirectory() as temporary:
        folder = Path(temporary)

        def write(name, values, sheet_name=HISTORY_SHEET, columns=HISTORY_COLUMNS):
            path = folder / name
            book = Workbook()
            sheet = book.active
            sheet.title = sheet_name
            sheet.append(columns)
            for value in values:
                sheet.append([value.get(key) for key in columns])
            book.save(path)
            book.close()
            return path

        base = {'CS NO': 'CS-MULTI-1', '현상': '펌프 누수', '조치': '씰 교체 후 정상화',
                'SITE': 'CN-A', '불량발생년월': '2026-09', '인원': 0}
        variants = [base, *(base | {key: f'구분 {index}'} for index, key in enumerate(HISTORY_DETAIL_COLUMNS)),
                    base | {'현상': '운전 중 펌프 누수'}, base | {'조치': '씰 발주, 미교체'},
                    base | {'인원': '0'}]
        path = write('multiple.xlsx', [base, *variants])
        rows = records.read_excel(path, 'CSList', **options)
        assert len(rows) == len(variants) == 11  # 12행 중 완전 중복 한 행만 제외
        assert len({records.record_id(row) for row in rows}) == 11
        assert rows[0]['_cell_locations']['현상'].endswith('!T2')
        assert rows[0]['SITE'] == 'CN-A' and rows[0]['불량발생년월'] == '2026-09'
        moved = write('reordered.xlsx', list(reversed(variants)), columns=HISTORY_COLUMNS[::-1])
        reordered = records.read_excel(moved, 'CSList', **options)
        assert {records.record_id(row) for row in rows} == {records.record_id(row) for row in reordered}
        cache = folder / 'multiple.rag.jsonl.gz'
        assert source_cache.convert_excel(path, cache, **options)['rows'] == 11
        assert records.read_excel(cache, 'CSList', **options) == rows
        # 정상 완료처럼 보이는 손상 캐시의 중복 상세와 잘못된 플래그는 거절합니다.
        contents = [json.loads(line) for line in gzip.decompress(cache.read_bytes()).decode().splitlines()]
        duplicate = copy.deepcopy(contents)
        duplicate[2]['values'], duplicate[2]['typed'] = duplicate[1]['values'], duplicate[1]['typed']
        bad_flag = copy.deepcopy(contents)
        bad_flag[0]['history_details'] = 'true'
        for bad in (duplicate, bad_flag):
            broken = folder / 'broken.rag.jsonl.gz'
            broken.write_bytes(gzip.compress(('\n'.join(json.dumps(item) for item in bad) + '\n').encode()))
            try:
                records.read_excel(broken, 'CSList', **options)
            except ValueError:
                pass
            else:
                raise AssertionError('손상된 상세 캐시를 허용했습니다.')
        cfg = folder / 'config.json'
        cfg.write_text(json.dumps({'model': 'mock', 'base_url': 'http://127.0.0.1:1/v1', 'api_key': 'EMPTY',
                                   'reasoning_effort': 'high', 'timeout_seconds': 1}))
        app = web.App(root, cfg, folder / 'out', root / 'cslist_schema.json')
        result = app.import_file(path, 'openpyxl')
        assert result['import']['added'] == 11
        listed = app.records({})
        assert listed['total'] == 11
        assert all(item['case_number'] == base['CS NO'] and item['customer'] == 'CN-A' for item in listed['items'])
        assert any('구분 0' in item['detail'] for item in listed['items'])
        detail = app.record({'id': [records.record_id(rows[0])]})
        assert {'SITE', '불량발생년월'} <= {cell['column'] for group in detail['groups'] for cell in group['cells']}
        assert app.import_file(moved, 'openpyxl')['import']['unchanged'] == 11
        with closing(app.connect()) as db:
            hits = search.retrieve(db, 'CS NO=CS-MULTI-1', limit=100)
            assert len(hits) == 11
            assert search.count_records(db, '전체 접수는 몇 건인가요?')['count'] == 1
            assert len({h['evidence'][0]['locations']['현상'] for h in hits}) == 11
            section = hits[0]['evidence'][0]
            packed = web.pack_answer('조치는?', '', hits, None, {'claims': [{'text': '과거 조치',
                'evidence': [{'id': section['id'], 'column': '조치', 'quote': section['cells']['조치']}]}]}, 0)
            cited = packed['claims'][0]['evidence'][0]
            assert cited['receipt'] == base['CS NO'] and cited['id'] == section['id']
            assert cited['location'] == section['locations']['조치']
            target = rows[2]
            before = dict(db.execute('SELECT id,original_json FROM normalized_records'))
            storage.apply_reviews(db, {'format': 'cell-reviews-v1', 'entries': [{
                'source_key': target['_source_key'], 'id': records.record_id(target),
                'revision': records.row_revision(target), 'column': '현상', 'action': 'edit',
                'value': '펌프에서 누수가 발생했습니다.'}]}, rules, schema=schema)
            assert dict(db.execute('SELECT id,original_json FROM normalized_records')) == before
            assert db.execute('SELECT record_id FROM cell_reviews').fetchall() == [(records.record_id(target),)]
            # 내용 수정은 동일 CS의 새 상세로 추가하며 기존 상세·검토를 지우지 않습니다.
            changed = write('changed.xlsx', [base | {'원인': '점검 후 확인'}])
            storage.ingest(db, records.read_excel(changed, 'CSList', **options), rules, schema=schema, append=True)
            assert db.execute('SELECT count(*) FROM normalized_records').fetchone()[0] == 12
            assert search.count_records(db, '전체 접수는 몇 건인가요?')['count'] == 1

        old_values = {'CS NO': 'CS-LEGACY-1', '현상': '펌프 누수', 'Site': 'CN-A', '불량발생월': '2026-09'}
        old_path = write('old.xlsx', [old_values], LEGACY_HISTORY_SHEET, LEGACY_HISTORY_COLUMNS)
        old_cache = folder / 'old.rag.jsonl.gz'
        source_cache.convert_excel(old_path, old_cache, **options)
        old_items = [json.loads(line) for line in gzip.decompress(old_cache.read_bytes()).decode().splitlines()]
        old_items[0].pop('history_details')  # 실제 v1 캐시에 없던 메타데이터
        old_cache.write_bytes(gzip.compress(('\n'.join(json.dumps(item) for item in old_items) + '\n').encode()))
        old_row = records.read_excel(old_cache, 'CSList', **options)[0]
        assert records.record_id(old_row) == '누적28:CS-LEGACY-1'
        assert row_schema(old_row, schema)['columns'] == LEGACY_HISTORY_COLUMNS
        new_path = write('new.xlsx', [old_values | {'SITE': 'CN-A', '불량발생년월': '2026-09'}])
        new_row = records.read_excel(new_path, 'CSList', **options)[0]
        assert records.history_content_key(old_row) == records.history_content_key(new_row)
        with closing(sqlite3.connect(':memory:')) as db:
            storage.ingest(db, [old_row], rules, schema=schema, append=True)
            storage.apply_reviews(db, {'format': 'cell-reviews-v1', 'entries': [{
                'source_key': old_row['_source_key'], 'id': records.record_id(old_row),
                'revision': records.row_revision(old_row), 'column': '현상', 'action': 'edit',
                'value': '펌프에서 누수가 발생했습니다.'}]}, rules, schema=schema)
            before = list(db.iterdump())
            assert storage.import_summary(db, [new_row])['unchanged'] == 1
            storage.ingest(db, [new_row], rules, schema=schema, append=True)
            assert list(db.iterdump()) == before, '구형 DB·검토·인용 출처가 바뀌었습니다.'
            old_hits = search.retrieve(db, 'CS NO=CS-LEGACY-1')
            assert english.translation_cells(json.dumps(old_hits[0]['evidence']), schema)['현상'] == '펌프에서 누수가 발생했습니다.'
        # 신형 상세를 먼저 등록한 뒤 구형 열 이름이나 다른 행 위치로 읽어도 재사용합니다.
        with closing(sqlite3.connect(':memory:')) as db:
            storage.ingest(db, [new_row], rules, schema=schema, append=True)
            before = list(db.iterdump())
            fresh_old = records.read_excel(old_path, 'CSList', **options)[0]
            moved_row = copy.deepcopy(new_row)
            moved_row['_excel_location'] = moved_row['_excel_location'].replace('2:', '3:')
            for incoming in (fresh_old, moved_row):
                assert storage.import_summary(db, [incoming])['unchanged'] == 1
                storage.ingest(db, [incoming], rules, schema=schema, append=True)
                assert list(db.iterdump()) == before
        assert app.read_rows(old_path, 'openpyxl')
        assert app.state()['schema_columns'] == LEGACY_HISTORY_COLUMNS
        # 신·구 시트가 같이 있으면 최신 시트 우선, 명시한 구형 시트는 그대로 선택합니다.
        from ragkit.input_formats import select_sheet
        assert select_sheet([LEGACY_HISTORY_SHEET, HISTORY_SHEET], 'CSList') == HISTORY_SHEET
        assert select_sheet([LEGACY_HISTORY_SHEET, HISTORY_SHEET], LEGACY_HISTORY_SHEET) == LEGACY_HISTORY_SHEET
    print('PASS: 부품별 상세·내용 차이·완전 중복·재정렬·캐시·CS 집계·구형 DB 및 검토 보존')


def main(root):
    check_details(root)
    schema, rules = cslist.load_schema(root / 'cslist_schema.json')
    options = {'required_columns': schema['columns'], 'id_column': schema['id_column']}
    config = {'base_url': 'http://127.0.0.1:1/v1', 'api_key': 'EMPTY', 'model': 'mock',
              'reasoning_effort': 'high', 'normalization_workers': 1, 'timeout_seconds': 1,
              'embedding': {'base_url': 'http://127.0.0.1:1/v1', 'model': 'mock', 'auto_index': False}}
    requests = []
    translations = {'冷却泵漏水': '냉각 펌프 누수', '运行2小时后漏水': '2小时 운전 후 누수가 발생했습니다.',
                    '更换密封后恢复正常': '씰 교체 후 정상화되었습니다.',
                    'Seal ordered; not replaced': '씰을 발주했으며, 교체는 하지 않았습니다.',
                    '原因待确认': '원인은 확인 대기 중입니다.', '颗粒增加': '파티클이 증가했습니다.'}

    def create(**request):
        requests.append(request)
        data = json.loads(request['messages'][1]['content'])
        assert '한국어로 번역' in request['messages'][0]['content']
        assert set(data['수정할 열']) <= set(HISTORY_REWRITE)
        content = {key: translations.get(data['원본 셀'][key], data['원본 셀'][key]) for key in data['수정할 열']}
        return NS(choices=[NS(finish_reason='stop', message=NS(content=json.dumps(content, ensure_ascii=False)))])

    embedded = []
    def embed(**kwargs):
        embedded.extend(kwargs['input'])
        return NS(data=[NS(index=i, embedding=[1., .1]) for i, _ in enumerate(kwargs['input'])])
    client = NS(chat=NS(completions=NS(create=create)), embeddings=NS(create=embed), close=lambda: None)

    def rejected(fn, message):
        try:
            fn()
        except ValueError as error:
            assert message in str(error), error
        else:
            raise AssertionError('거절해야 하는 입력: ' + message)

    with TemporaryDirectory() as temporary:
        folder = Path(temporary)
        path, cache = folder / 'history.xlsx', folder / 'history.rag.jsonl.gz'
        cfg = folder / 'config.json'
        cfg.write_text(json.dumps(config), encoding='utf-8')
        app = web.App(root, cfg, folder / 'output', root / 'cslist_schema.json')
        book = Workbook()
        sheet = book.active
        sheet.title = HISTORY_SHEET
        # 순서와 주변 공백은 달라도 원본 셀 위치를 정확히 유지합니다.
        columns = HISTORY_COLUMNS[::-1]
        sheet.append([' ' + c + ' ' for c in columns])
        samples = [
            {'CS NO': 'CS-DEMO-001', '제목': '冷却泵漏水', '현상': '运行2小时后漏水',
             '원인': '', '조치': '更换密封后恢复正常', '설비': 'SCW132', '인원': 0,
             '불량발생일': datetime(2026, 8, 31), 'SITE': 'CN-A'},
            {'CS NO': 'CS-NEW-002', '제목': 'Leak check', '원인': '原因待确认',
             '현상': '运行2小时后漏水', '조치': 'Seal ordered; not replaced'},
            {'CS NO': 12345, '현상': '颗粒增加', 'Defect 상세': '颗粒增加'},
        ]
        for values in samples:
            sheet.append([values.get(key) for key in columns])
        book.save(path)
        incoming = app.read_rows(path, 'openpyxl')
        assert len(incoming) == 3 and incoming[0]['CS NO'] == 'CS-DEMO-001'
        identities = [records.record_id(row) for row in incoming]
        assert len(set(identities)) == 3 and all(i.startswith('누적28상세:') for i in identities)
        assert incoming[0]['SITE'] == 'CN-A'
        assert incoming[0]['원인'] == '' and incoming[0]['인원'] == '0'
        assert incoming[0]['_cell_locations']['CS NO'].endswith('!AB2')
        assert incoming[0]['_typed_cells']['불량발생일']['kind'] == 'datetime'
        assert records.read_excel(path, HISTORY_SHEET, **options) == incoming
        inspected = app.inspect({'path': str(path)})
        assert inspected['columns'] == 28 and inspected['id_column'] == 'CS NO'
        assert app.state()['columns'] == 28 and app.state()['sheet'] == HISTORY_SHEET
        source_cache.convert_excel(path, cache, **options)
        assert records.read_excel(cache, 'CSList', **options) == incoming

        # Windows COM 경계만 대체합니다. 새 시트 선택·공통 검사·변환은 실제 코드입니다.
        def excel_range(start, end=None):
            if end is None:
                cell = sheet.cell(*start)
                value = to_excel(cell.value) if isinstance(cell.value, datetime) else cell.value
                return NS(api=NS(Value2=value), number_format=cell.number_format)
            block = NS(value=[[c.value for c in row] for row in sheet.iter_rows()],
                       api=NS(MergeCells=False, HasFormula=False))
            block.options = Mock(return_value=block)
            return block
        fake_sheet = NS(name=HISTORY_SHEET, range=excel_range,
                        api=NS(Cells=NS(Find=lambda **kw: NS(Row=4, Column=28))))
        fake_book = NS(fullname=str(path), sheets=[fake_sheet], api=NS(Date1904=False))
        with patch.dict(sys.modules, xlwings=NS(apps=[NS(books=[fake_book])])), patch('sys.platform', 'win32'):
            assert records.read_excel(path, 'CSList', reader='xlwings', **options) == incoming
            source_cache.convert_excel(path, cache, reader='xlwings', **options)
        assert records.read_excel(cache, 'CSList', **options) == incoming
        # 누락 열은 거절하고 같은 CS NO의 다른 내용은 상세로 보존합니다.
        sheet.cell(1, 1, '없는 열')
        book.save(path)
        rejected(lambda: app.read_rows(path, 'openpyxl'), '첫 행의 열')
        sheet.cell(1, 1, columns[0])
        sheet.cell(3, 28, 'CS-DEMO-001')
        book.save(path)
        duplicate_number = app.read_rows(path, 'openpyxl')
        assert len(duplicate_number) == 3
        assert len({records.record_id(row) for row in duplicate_number}) == 3
        with patch.dict(sys.modules, xlwings=NS(apps=[NS(books=[fake_book])])), patch('sys.platform', 'win32'):
            assert records.read_excel(path, 'CSList', reader='xlwings', **options) == duplicate_number
        book.close()
        path.unlink()  # 이후 모든 단계는 DB 원문으로만 수행합니다.

        legacy = records.read_excel(root / 'cslist_sample.xlsx', 'CSList', **options)[:1]
        with closing(sqlite3.connect(app.db_path)) as db:
            storage.ingest(db, legacy, rules, schema=schema, append=True)
            before = db.execute('SELECT * FROM normalized_records').fetchone()
            storage.ingest(db, incoming, rules, schema=schema, append=True)
            assert db.execute('SELECT * FROM normalized_records WHERE id=?', ('CS-DEMO-001',)).fetchone() == before
            assert search.count_records(db, '전체 접수는 몇 건인가요?')['count'] == 3
            assert search.count_records(db, 'CS NO="CS-DEMO-001" 접수는 몇 건인가요?')['count'] == 1
            exact = search.retrieve(db, 'CS-DEMO-001', limit=10)
            assert {h['record_id'] for h in exact} == {'CS-DEMO-001', identities[0]}
            assert search.retrieve(db, 'CS NO=12345')[0]['record_id'] == identities[2]
            assert not search.retrieve(db, 'CS NO=98765')
            with patch.object(semantic, 'vector_hits', side_effect=AssertionError('번호 검색을 벡터로 확장')):
                assert semantic.search_cases(db, 'CS NO=12345', config)[0]['record_id'] == identities[2]
                assert not semantic.search_cases(db, 'CS NO=98765', config)
            assert not english.bilingual_retrieve(db, 'CS NO=98765', None, config)
            assert not search.retrieve(db, '파티클'), '미번역 제안을 검색했습니다.'
            stored = [json.loads(raw) for (raw,) in db.execute("SELECT original_json FROM normalized_records WHERE id LIKE '누적28상세:%'")]
            prepared = storage.ingest(db, stored, rules, client, config, schema=schema, append=True)
            assert len(requests) == 3
            assert next(r for r in prepared if r['id'] == identities[0])['normalized']['원인'] == ''
            assert not search.retrieve(db, '파티클'), '미채택 번역을 검색했습니다.'
            accepted = storage.accept_pending_reviews(db, rules, schema)
            assert not accepted['failures'] and accepted['applied_rows'] == 3, accepted
            assert search.retrieve(db, '파티클')[0]['record_id'] == identities[2]
            hits = search.retrieve(db, 'CS-DEMO-001', limit=10)
            translated = next(h for h in hits if h['record_id'].startswith('누적28상세:'))
            evidence = translated['evidence'][0]
            assert evidence['cells']['조치'] == '更换密封后恢复正常'
            assert evidence['accepted']['조치'] == translations['更换密封后恢复正常']
            assert evidence['locations']['현상'] == incoming[0]['_cell_locations']['현상']
            en = english.translation_cells(json.dumps(translated['evidence']), schema)
            assert en['조치'] == translations['更换密封后恢复正常'] and 'CS NO' not in en
            storage.ingest(db, incoming, rules, client, config, schema=schema, append=True)
            assert len(requests) == 3, '같은 파일을 다시 번역했습니다.'
            assert semantic.build_vectors(db, client, config)['generated'] >= 3
            assert any('파티클' in text for text in embedded)
            assert any('SCW132' in text for text in embedded)
            changed = copy.deepcopy(incoming)
            changed[0]['_source_key'] = '0' * 64
            changed[0]['현상'] = '새 원문'
            changed_summary = {}
            storage.ingest(db, changed, rules, schema=schema, append=True, import_result=changed_summary)
            assert changed_summary['added'] == 1 and changed_summary['unchanged'] == 2
            # 한국어 답변에도 중국어 원문 전체가 인용되고, 없는 증상·단위는 거절합니다.
            def claim(text):
                return {'claims': [{'text': text, 'evidence': [{'id': evidence['id'], 'column': key} for key in ('현상', '조치')]}]}
            payload = claim('2小时 운전 후 누수가 발생했고 씰 교체 후 정상화되었습니다.')
            assert '정상화' in llm.validate_claims(payload, [translated])
            assert payload['claims'][0]['evidence'][0]['quote'] == incoming[0]['현상']
            rejected(lambda: llm.validate_claims(claim('진동이 발생했습니다.'), [translated]), '증상')
            rejected(lambda: llm.validate_claims(claim('2시간 운전했습니다.'), [translated]), '단위')
            numbered = search.retrieve(db, 'CS NO=12345')
            def numeric_claim(text):
                return {'claims': [{'text': text, 'evidence': [{'id': numbered[0]['evidence'][0]['id'], 'column': '현상'}]}]}
            rejected(lambda: llm.validate_claims(numeric_claim('측정값은 12345입니다.'), numbered), '숫자')
            assert '12345' in llm.validate_claims(numeric_claim('CS NO 12345에서 파티클이 증가했습니다.'), numbered)
            for key, value, message in [('현상', '3小时 운전 후 누수가 발생했습니다.', '숫자'),
                                        ('현상', '2分钟 운전 후 누수가 발생했습니다.', '단위'),
                                        ('원인', '원인이 없습니다.', '빈 셀'), ('인원', '1', '보존'),
                                        ('설비', 'OTHER132', '식별자')]:
                rejected(lambda: llm.validate_normalized(records.public_cells(incoming[0]) | {key: value}, incoming[0], rules), message)
        detail = app.record({'id': [identities[0]]})
        fields = [cell for group in detail['groups'] for cell in group['cells']]
        assert len(fields) == 28
        assert next(cell for cell in fields if cell['column'] == '조치')['editable']
        assert next(cell for cell in fields if cell['column'] == '현상')['proposal'] == translations['运行2小时后漏水']
        assert app.records({})['total'] == 5
        # 저장 자료 CLI에서 기존 id_column을 새 자료에도 강제하지 않습니다.
        cli.main(root, ['--schema', str(root / 'cslist_schema.json'), '--config', str(cfg),
                        '--output', str(app.output), '--saved', '--question', '전체 접수는 몇 건인가요?'])
    print('PASS: 누적 CS 28열·다국어 번역·두 양식 보존·번호 통합 조회·검토·벡터·원문 인용')


if __name__ == '__main__':
    main(Path(__file__).resolve().parents[1] / 'windows-demo')
