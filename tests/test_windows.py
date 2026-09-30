"""서버 호출 없이 Excel·검색·요청 옵션을 검증합니다."""
import json
from datetime import datetime
import sqlite3
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from ragkit import records, llm, storage, search, report as reports, cli, evaluation


def check_field015_response(row, rules, config):
    # 외부 서버의 응답만 대체합니다. 실제 Excel 행과 정리·검증 코드를 그대로 실행합니다.
    def response(value):
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop',
            message=SimpleNamespace(content=json.dumps(value, ensure_ascii=False)))])

    proposal = {'인수인계': '9/8 야간 인수인계에서는 이상음 재확인을 요청했습니다. 대상·측정값·조치 이력은 기재되지 않았습니다.'}
    original = records.public_cells(row)
    create = Mock()
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    # 수정 대상만 반환하거나, 보존 셀을 원문 그대로 함께 반환하는 경우 모두 처리합니다.
    for result in (proposal, original | proposal):
        create.reset_mock(side_effect=True)
        create.return_value = response(result)
        actual = llm.normalize(client, config, row, rules)
        assert actual == original | proposal
        assert actual['대상 장비'] == actual['접수 메모'] == actual['작업자 기록'] == ''
        assert create.call_count == 1

    # 대상 열 누락 시 한 번 보정을 요청하며, 누락을 원문 복사로 숨기지 않습니다.
    create.reset_mock()
    create.side_effect = [response({}), response(proposal)]
    assert llm.normalize(client, config, row, rules) == original | proposal
    assert create.call_count == 2
    correction = create.call_args.kwargs['messages'][-1]['content']
    assert '인수인계' in correction and '누락' in correction

    # 빈칸 채우기·ID 변경은 중단합니다. 숫자 오류는 해당 셀을 원문으로 복귀시킵니다.
    for invalid, detail in [({}, '인수인계'), (proposal | {'요약': '추가 열'}, '요약'),
                            (proposal | {'대상 장비': '알 수 없음'}, '빈 셀'),
                            (proposal | {'id': 'FIELD-014'}, 'ID'),
                            ({'인수인계': proposal['인수인계'].replace('9/8', '9/9')}, '숫자')]:
        create.reset_mock(side_effect=True)
        create.return_value = response(invalid)
        try:
            result = llm.normalize(client, config, row, rules)
        except ValueError as error:
            assert 'FIELD-015' in str(error) and detail in str(error)
        else:
            assert detail == '숫자' and result == original
            assert set(result.fallbacks) == {'인수인계'}
            llm.validate_normalized(result, row, rules)
        assert create.call_count == 2

    # 연결 장애는 형식 보정으로 해결할 수 없으므로 추가 요청하지 않습니다.
    create.reset_mock()
    create.side_effect = TimeoutError('서버 응답 지연')
    try:
        llm.normalize(client, config, row, rules)
    except TimeoutError:
        pass
    else:
        raise AssertionError('서버 장애를 숨겼습니다.')
    assert create.call_count == 1
    print('PASS: FIELD-015 보존 셀 동봉·누락 보정·재시도 한도·빈칸/ID/숫자 보호')


def main(root):
    rules = llm.load_rules(root / "column_rules.json")
    rows = records.read_excel(root / 'sample_data.xlsx')
    assert len(rows) == 6 and rows[0]['id'] == 'DEMO-001'
    assert rows[0]['_excel_location'].endswith('점검기록!A2:H2')
    # 잘못된 입력은 실제 파일을 저장하지 않고 메모리 복사본에서 검사합니다.
    for cell, value, message in [('A1', '문서명', '첫 행'), ('A3', 'DEMO-001', '중복 ID'), ('F2', '=1+1', '수식')]:
        book = load_workbook(root / 'sample_data.xlsx')
        book['점검기록'][cell] = value
        with patch.object(records, 'load_workbook', return_value=book):
            try:
                records.read_excel(root / 'sample_data.xlsx')
            except ValueError as error:
                assert message in str(error)
            else:
                raise AssertionError('잘못된 Excel 입력을 허용했습니다.')
    assert search.terms('냉각수를') & search.terms('냉각수가') == {'냉각', '각수'}
    with sqlite3.connect(':memory:') as db:
        storage.ingest(db, rows, rules)
        question = '냉각기 A의 냉각수를 보충했는데도 유량이 낮습니다. 무엇을 확인해야 하나요?'
        hits = search.retrieve(db, question)
        assert 'DEMO-002' in {h['id'] for h in hits}
        assert search.retrieve(db, '컨베이어 C에서 제품을 감지하지 못할 때 어떻게 조치하나요?')[0]['id'] == 'DEMO-004'
        try:
            search.retrieve(db, '!!!')
        except ValueError:
            pass
        else:
            raise AssertionError('빈 검색어를 허용했습니다.')
        storage.ingest(db, rows[:1], rules)
        assert db.execute('select count(*) from normalized_records').fetchone()[0] == 1
    # 실제 Excel의 빈 셀·완전 빈 행·자유로운 열 이름을 검사합니다.
    realistic = records.read_excel(root / 'realistic_data.xlsx')
    assert len(realistic) == 15
    by_id = {row['id']: row for row in realistic}
    assert by_id['FIELD-013']['작업자 기록'] == ''
    assert by_id['FIELD-014']['대상 장비'] == ''
    assert by_id['FIELD-015']['접수 메모'] == ''
    assert by_id['FIELD-015']['_excel_location'].endswith('A17:E17')
    assert '이상음' in records.text_for(by_id['FIELD-015'])
    book = load_workbook(root / 'realistic_data.xlsx')
    book['점검기록']['A2'] = None
    with patch.object(records, 'load_workbook', return_value=book):
        try:
            records.read_excel(root / 'realistic_data.xlsx')
        except ValueError as error:
            assert 'id' in str(error)
        else:
            raise AssertionError('내용이 있는데 ID 없는 행을 허용했습니다.')
    checks = json.loads((root / 'realistic_cases.json').read_text(encoding='utf-8'))
    with sqlite3.connect(':memory:') as db:
        storage.ingest(db, realistic, rules)
        for check in checks:
            found = {h['id'] for h in search.retrieve(db, check['question'])}
            assert set(check['expected_ids']).issubset(found), (check['question'], found)
    print(f'PASS: 빈 셀·빈 행·ID 누락 검사 및 현실형 검색 {len(checks)}건')
    config = json.loads((root / 'config.json').read_text(encoding='utf-8'))
    assert config['base_url'] == 'http://23.43.51.216:6100/v1'
    assert config['api_key'] == 'EMPTY'
    assert config['model'] == 'gpt-4' and config['reasoning_effort'] == 'high'
    expected_options = {'temperature': 1, 'top_p': 1, 'max_tokens': 131072}
    assert config['context_window'] == 131072
    assert llm.generation_options(config) == expected_options
    assert llm.generation_options(config, for_answer=True) == expected_options
    assert llm.generation_options(config | {'answer_max_tokens': 8192}, for_answer=True) == expected_options | {'max_tokens': 8192}
    messages = [{'role': 'user', 'content': '압력 -0.2 MPa 점검. ' * 100}]
    reserved = len(json.dumps(messages, ensure_ascii=False).encode('utf-8')) + 1024
    assert llm.generation_options(config, messages=messages)['max_tokens'] == 131072 - reserved
    longer = messages + [{'role': 'user', 'content': '숫자를 보존해 다시 답하세요.'}]
    assert llm.generation_options(config, messages=longer)['max_tokens'] < 131072 - reserved
    assert llm.generation_options(config | {'answer_max_tokens': 8192}, for_answer=True, messages=messages)['max_tokens'] == 8192
    assert 'max_tokens' not in llm.generation_options({'reasoning_effort': 'high'}, messages=messages)
    try:
        llm.generation_options(config | {'context_window': 2048}, messages=messages)
    except ValueError as error:
        assert '입력 예약량' in str(error)
    else:
        raise AssertionError('문맥 한도를 넘는 요청을 허용했습니다.')
    for invalid in ({'max_tokens': True}, {'max_tokens': 0}, {'temperature': float('nan')}, {'top_p': 1.1}, {'top_p': False}, {'context_window': True}, {'context_window': 0}):
        try:
            llm.generation_options(config | invalid)
        except ValueError:
            pass
        else:
            raise AssertionError('잘못된 생성 옵션을 허용했습니다.')
    check_field015_response(by_id['FIELD-015'], rules, config)
    # 셀별 수정문은 원래 열 구조와 빈 값을 유지해야 합니다.
    messy = {'id': 'NOTE-001', '원인': 'filter 막힘?', '조치': '청소함', '결과': '',
             '측정': '7.2 → 11.8 L/min', '비고': '  '}
    normalized = dict(messy, 원인='필터(여과망) 막힘이 의심됩니다.', 조치='청소를 실시했습니다.')
    normal_create = Mock(return_value=SimpleNamespace(choices=[SimpleNamespace(
        finish_reason='stop', message=SimpleNamespace(content=json.dumps({k:v for k,v in normalized.items() if llm.editable(messy, k, rules)}, ensure_ascii=False)))]))
    normal_client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=normal_create)))
    with sqlite3.connect(':memory:') as db:
        prepared = storage.ingest(db, [messy], rules, normal_client, config)
        storage.ingest(db, [messy], rules, normal_client, config)
        assert normal_create.call_count == 1
        assert search.retrieve(db, '여과망') == []  # 미채택 제안은 검색되지 않습니다.
        entry = dict(source_key='', id=messy['id'], revision=records.row_revision(messy), column='원인', action='accept', value=normalized['원인'])
        storage.apply_reviews(db, {'format': 'cell-reviews-v1', 'entries': [entry]}, rules)
        assert search.retrieve(db, '여과망')[0]['id'] == 'NOTE-001'
        # 잘못된 항목 하나가 섞여도 앞의 올바른 항목만 부분 반영되면 안 됩니다.
        saved_reviews = db.execute('SELECT * FROM cell_reviews').fetchall()
        saved_search = db.execute('SELECT search_text FROM normalized_records').fetchall()
        for invalid_entry in (entry, dict(entry, source_key='다른 파일'), dict(entry, column='측정')):
            try:
                storage.apply_reviews(db, {'format': 'cell-reviews-v1', 'entries': [dict(entry, action='original'), invalid_entry]}, rules)
            except ValueError:
                pass
            else:
                raise AssertionError('중복·다른 파일·보존 열의 검토를 허용했습니다.')
            assert db.execute('SELECT * FROM cell_reviews').fetchall() == saved_reviews
            assert db.execute('SELECT search_text FROM normalized_records').fetchall() == saved_search
        payload = normal_create.call_args.kwargs
        assert json.loads(payload['messages'][1]['content'])['원본 셀'] == messy
        assert '측정' not in json.loads(payload['messages'][1]['content'])['수정할 열']
        # 직접 편집을 저장한 뒤 제안을 다시 생성해도 채택문은 유지됩니다.
        entry.update(action='edit', value='여과망이 막혔을 가능성이 있습니다.')
        storage.apply_reviews(db, {'format': 'cell-reviews-v1', 'entries': [entry]}, rules)
        again = storage.ingest(db, [messy], rules, normal_client, config)
        assert again[0]['effective']['원인'] == entry['value']
        assert normal_create.call_count == 1
        stored = json.loads(db.execute('SELECT normalized_json FROM normalized_records').fetchone()[0])
        assert stored == normalized and list(stored) == list(messy)
        assert normal_create.call_args.kwargs['extra_body'] == {'reasoning_effort': 'high'}
        assert {key: normal_create.call_args.kwargs[key] for key in ('temperature', 'top_p')} == {'temperature': 1, 'top_p': 1}
        request = normal_create.call_args.kwargs
        assert 0 < request['max_tokens'] < config['context_window']
        assert request['max_tokens'] + len(json.dumps(request['messages'], ensure_ascii=False).encode('utf-8')) + 1024 <= config['context_window']
        assert 'max_completion_tokens' not in normal_create.call_args.kwargs
        changed = dict(messy, 조치='청소함. 재점검 예정')
        new = storage.ingest(db, [changed], rules, normal_client, config)
        assert new[0]['reviews']['원인']['stale'] and new[0]['effective']['원인'] == messy['원인']
        assert new[0]['reviews']['원인']['previous'] == entry['value']
        try:
            storage.apply_reviews(db, {'format': 'cell-reviews-v1', 'entries': [entry]}, rules)
        except ValueError:
            pass
        else:
            raise AssertionError('원문이 변경된 이전 검토 파일을 적용했습니다.')
        assert normal_create.call_count == 2
        previous = db.execute('SELECT * FROM normalized_records').fetchall()
        # 두 번째 행이 실패할 때 첫 번째 행만 부분 저장되지 않아야 합니다.
        failed_response = SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content='{}'))])
        normal_create.side_effect = [normal_create.return_value, failed_response, failed_response]
        try:
            storage.ingest(db, [messy, dict(messy, id='NOTE-002')], rules, normal_client, config)
        except ValueError:
            pass
        else:
            raise AssertionError('잘못된 셀 응답을 허용했습니다.')
        assert db.execute('SELECT * FROM normalized_records').fetchall() == previous
    invalid = [dict(normalized, id='다른 ID'), dict(normalized, 결과='정상'),
               dict(normalized, 비고=''), dict(normalized, 조치=''),
               dict(normalized, 측정='7.2 → 12 L/min'), dict(normalized, 측정=None),
               dict(normalized, 요약='새 열'), {k:v for k,v in normalized.items() if k != '원인'}]
    for result in invalid:
        try:
            llm.validate_normalized(result, messy, rules)
        except ValueError:
            pass
        else:
            raise AssertionError(f'잘못된 수정문을 허용했습니다: {result}')
    # 실제 형식의 숫자·날짜·참거짓 값은 원본 자료형 정보와 함께 보존합니다.
    book = load_workbook(root / 'sample_data.xlsx')
    book['점검기록']['D2'] = 0.18
    book['점검기록']['E2'] = datetime(2026, 9, 9, 10, 30)
    book['점검기록']['F2'] = False
    with patch.object(records, 'load_workbook', return_value=book):
        typed = records.read_excel(root / 'sample_data.xlsx')[0]
    assert typed['_typed_cells']['증상']['value'] == 0.18
    assert typed['_typed_cells']['기록된 원인']['kind'] == 'datetime'
    assert typed['_typed_cells']['조치 방법']['value'] is False
    assert not llm.editable(typed, '증상', rules) and typed['조치 방법'] == 'False'
    assert not llm.editable({'id':'X', '알수없는열':'메모'}, '알수없는열', rules)
    # 누락 셀이 있는 실제 샘플도 열과 빈 값을 그대로 보존합니다.
    for row in realistic:
        result = {k:v for k,v in row.items() if not k.startswith('_')}
        llm.validate_normalized(result, row, rules)
    assert by_id['FIELD-015']['_cell_locations']['인수인계'].endswith('!E17')
    # 네트워크를 사용하는 클라이언트 대신 요청 내용을 기록하는 대역을 사용합니다.
    create = Mock(return_value=SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content='검사용 모의 답변 [DEMO-002]'))]))
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    reply, context = llm.answer(client, config, question, hits)
    payload = create.call_args.kwargs
    assert payload['model'] == 'gpt-4'
    assert payload['extra_body'] == {'reasoning_effort': 'high'}
    assert {key: payload[key] for key in ('temperature', 'top_p')} == {'temperature': 1, 'top_p': 1}
    assert config['answer_tokenizer'] == 'o200k_harmony'
    input_tokens = sum(len(llm._answer_encoder().encode_ordinary(m['content'])) for m in payload['messages'])
    assert payload['max_tokens'] + input_tokens + 1024 == config['context_window']
    assert 'max_completion_tokens' not in payload
    assert context in payload['messages'][1]['content']
    create.return_value.choices[0].finish_reason = 'length'
    try:
        llm.answer(client, config, question, hits)
    except RuntimeError:
        pass
    else:
        raise AssertionError('잘린 답변을 허용했습니다.')
    with tempfile.TemporaryDirectory() as tmp:
        report = dict(config, column_rules=rules, created_at='test', source_file='sample_data.xlsx', source_sheet='점검기록', rows=rows,
                      normalized_records=prepared, cases=[dict(question=question, hits=hits, answer='<script>mock</script>', seconds=0)])
        reports.render(report, Path(tmp) / 'demo.html')
        rendered = (Path(tmp) / 'demo.html').read_text(encoding='utf-8')
        assert '&lt;script&gt;mock&lt;/script&gt;' in rendered and '<script>mock' not in rendered
        assert 'gpt-4' in rendered and 'high' in rendered
        assert '<ins>' in rendered and '<del>' in rendered and '셀별 원문·수정문 비교' in rendered
        assert 'download-reviews' in rendered and '미채택' not in storage.search_content(prepared[0])
        assert '빈 셀 유지' in rendered
    # 실제 main을 임시 출력 폴더에서 실행합니다. 모의 응답은 품질 평가 결과가 아닙니다.
    def fake_create(**payload):
        if payload['messages'][0]['content'].startswith(llm.NORMAL_PROMPT):
            value = json.loads(payload['messages'][1]['content'])
            # 원문 셀을 그대로 반환하여 구조·빈칸·전체 데이터 흐름만 확인합니다.
            content = json.dumps({key:value['원본 셀'][key] for key in value['수정할 열']}, ensure_ascii=False)
        else:
            content = '검사용 모의 답변입니다.'
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content=content))])
    fake = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=fake_create)))
    source_root = root
    before = (source_root / 'realistic_data.xlsx').read_bytes()
    with tempfile.TemporaryDirectory() as tmp:
        output = Path(tmp)
        (output / 'column_rules.json').write_bytes((source_root / 'column_rules.json').read_bytes())
        (output / 'config.json').write_bytes((source_root / 'config.json').read_bytes())
        args = ['demo.py', '--excel', str(source_root / 'realistic_data.xlsx'),
                '--cases', str(source_root / 'realistic_cases.json')]
        with patch.object(cli, 'OpenAI', return_value=fake):
            cli.main(output, args[1:])
        report = json.loads((output / 'demo_results.json').read_text(encoding='utf-8'))
        assert len(report['normalized_records']) == 15 and len(report['cases']) == 14
        for record in report['normalized_records']:
            assert record['normalized'] == {k:v for k,v in record['original_cells'].items() if not k.startswith('_')}
        assert 'realistic_data.xlsx / 점검기록!E17' in (output / 'demo.html').read_text(encoding='utf-8')
        assert (source_root / 'realistic_data.xlsx').read_bytes() == before
    # 평가 자료의 참조 문장도 형식·빈칸·숫자 보존 기준을 만족하는지 확인합니다.
    dataset = json.loads((source_root / 'evaluation_cases.json').read_text(encoding='utf-8'))
    assert len(dataset['cases']) == 77 and len({c['id'] for c in dataset['cases']}) == 77
    for case in dataset['cases']:
        llm.validate_normalized(case['input'] | {'메모': case['reference']}, case['input'], rules)
    with tempfile.TemporaryDirectory() as tmp:
        evaluated = evaluation.run_evaluation(fake, config, dataset, Path(tmp), rules, limit=3)
        assert all('error' not in item for item in evaluated['results'])
        assert len(list(Path(tmp).glob('*.html'))) == 1
        assert all(item['proposal']['결과'] == '' for item in evaluated['results'])
        bad_case = dataset['cases'][0] | {'input': {'id': dataset['cases'][0]['id'], '다른 열': '메모 없음'}}
        try:
            evaluation.run_evaluation(None, config, dataset | {'cases': [bad_case]}, Path(tmp) / 'invalid', rules)
        except ValueError as error:
            assert '메모' in str(error)
        else:
            raise AssertionError('평가 자료에 필요한 메모 열 누락을 허용했습니다.')
        assert not (Path(tmp) / 'invalid').exists()
    print('PASS: 평가 자료 77건의 구조·참조문, 평가 실행 및 보고서')
    print('PASS: 셀별 JSON·열/ID/빈 셀/숫자 보존·변경 캐시·실패 시 DB 보존·Excel·DB 동기화·키워드 검색·gpt-4/high 요청·잘린 답변·HTML 이스케이프')
