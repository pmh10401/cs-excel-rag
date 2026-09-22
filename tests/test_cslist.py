"""CS 128열 입력부터 작업별 검색·검토 저장까지 서버 없이 확인합니다."""
import json
import sqlite3
import tempfile
import httpx2
from openai import APITimeoutError
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from openpyxl import load_workbook
from ragkit import records, llm, storage, search, cli, report


def main(root):
    schema = json.loads((root / 'cslist_schema.json').read_text(encoding='utf-8'))
    excel = root / 'cslist_sample.xlsx'
    before = excel.read_bytes()
    rows = records.read_excel(excel, None, required_columns=schema['columns'], id_column='CS 접수번호')
    assert len(rows) == 8 and list(records.public_cells(rows[0])) == schema['columns']
    assert len(records.public_cells(rows[0])) == 128 and 'id' not in rows[0]
    assert records.record_id(rows[0]) == 'CS-DEMO-001'
    assert rows[0]['_source_sheet'] == 'CSList'
    assert rows[0]['_typed_cells']['접수일시']['kind'] == 'datetime'
    assert rows[5]['무상자재금액'] == '0' and rows[5]['MTTR'] == ''
    # 열 위치가 달라도 이름으로 작업을 묶고 실제 셀 주소는 새 위치로 보존합니다.
    book = load_workbook(excel)
    for cells in book['CSList'].iter_rows():
        values = [cell.value for cell in cells][::-1]
        for cell, value in zip(cells, values):
            cell.value = value
    with patch.object(records, 'load_workbook', return_value=book):
        reordered = records.read_excel(excel, None, required_columns=schema['columns'], id_column='CS 접수번호')
    assert records.public_cells(reordered[0]) == records.public_cells(rows[0])
    assert reordered[0]['_cell_locations']['CS 접수번호'].endswith('!DX2')
    book = load_workbook(excel)
    book.create_sheet('추가 시트')
    with patch.object(records, 'load_workbook', return_value=book):
        try:
            records.read_excel(excel, None, required_columns=schema['columns'], id_column='CS 접수번호')
        except ValueError as error:
            assert '시트가 여러 개' in str(error)
        else:
            raise AssertionError('여러 시트를 임의로 선택했습니다.')
    # 중복 접수번호를 합치거나 행 번호로 임의 보정하면 안 됩니다.
    for column, value, expected in [('A3', 'CS-DEMO-001', '중복'), ('A2', '', 'CS 접수번호'), ('B1', '다른 제목', '제목')]:
        book = load_workbook(excel)
        book['CSList'][column] = value
        with patch.object(records, 'load_workbook', return_value=book):
            try:
                records.read_excel(excel, None, required_columns=schema['columns'], id_column='CS 접수번호')
            except ValueError as error:
                assert expected in str(error)
            else:
                raise AssertionError('잘못된 CS 입력을 허용했습니다.')
    from ragkit.cslist import load_schema
    schema, rules = load_schema(root / 'cslist_schema.json')
    config = json.loads((root / 'config.json').read_text(encoding='utf-8'))
    payloads = []
    def fake_create(**request):
        payloads.append(request)
        if request['messages'][0]['content'] == llm.NORMAL_PROMPT:
            data = json.loads(request['messages'][1]['content'])
            original = data['원본 셀']
            # 서로 다른 작업 내용을 같은 수정 요청에 넣지 않았는지 확인합니다.
            assert not (any(k.endswith('(1)') for k in original) and any(k.endswith('(2)') for k in original))
            value = {key: original[key] + ' (모의 수정문)' for key in data['수정할 열']}
            content = json.dumps(value, ensure_ascii=False)
        else:
            sources = json.loads(request['messages'][1]['content'].split('원문 자료:\n', 1)[1])
            source = sources[0]
            column = next(key for key, value in source['cells'].items() if value.strip() and key != 'CS 접수번호')
            content = json.dumps({'claims': [{'text': '기록에 적힌 내용입니다.', 'evidence': [
                {'id': source['id'], 'column': column, 'quote': source['cells'][column]}]}]}, ensure_ascii=False)
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content=content))])
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=fake_create)))
    with sqlite3.connect(':memory:') as db:
        prepared = storage.ingest(db, rows, rules, client, config, schema=schema)
        assert len(prepared[0]['normalized']) == 128
        assert prepared[0]['normalized']['CS 접수번호'] == 'CS-DEMO-001'
        assert prepared[0]['normalized']['접수일시'] == rows[0]['접수일시']
        assert prepared[4]['normalized']['원인(1)'] == ''
        assert not db.execute("SELECT 1 FROM sqlite_master WHERE name GLOB 'rag_chunks*'").fetchone()
        chunks = {section['id']: json.dumps(section, ensure_ascii=False)
                  for (encoded,) in db.execute('SELECT evidence FROM rag_documents') for section in json.loads(encoded)}
        assert 'CS-DEMO-001/작업1' in chunks and 'CS-DEMO-001/작업2' in chunks
        assert '배기 튜브' not in chunks['CS-DEMO-001/작업1']
        assert '영점 조정함' not in chunks['CS-DEMO-001/작업2']
        assert 'CS-DEMO-005/작업1' not in chunks
        assert '모의 수정문' not in ''.join(chunks.values())
        assert search.retrieve(db, '배기 튜브 균열 교체')[0]['id'] == 'CS-DEMO-001/문서'
        cases = json.loads((root / 'cslist_cases.json').read_text(encoding='utf-8'))
        for case in cases:
            assert search.retrieval_check(search.retrieve(db, case['question']), case) is True, case['question']
        assert json.loads(chunks['CS-DEMO-006/접수'])['cells']['무상자재금액'] == '0'
        # 문구가 비슷해도 명시한 번호·설비·날짜 밖의 기록을 섞지 않아야 합니다.
        assert search.retrieve(db, 'CS-DEMO-999의 조치 결과는?') == []
        exact = search.retrieve(db, 'cs-demo-001의 조치 결과는?')
        assert exact and {h['record_id'] for h in exact} == {'CS-DEMO-001'}
        dated = search.retrieve(db, 'DEMO-ETCH-01 2026-09-01 배기 튜브 교체 결과는?')
        assert 'CS-DEMO-001/작업2' in search.evidence_ids(dated)
        assert search.retrieve(db, 'DEMO-ETCH-01 2026-09-02 배기 튜브 교체 결과는?') == []
        entry = {'source_key':rows[0]['_source_key'], 'id':'CS-DEMO-001', 'revision':records.row_revision(rows[0]),
                 'column':'조치내용(2)', 'action':'edit', 'value':'배기 튜브 교체. 60분 관찰 중 알람 없음. 장기 재발 미확인. 검토확인표식.'}
        storage.apply_reviews(db, {'format':'cell-reviews-v1','entries':[entry]}, rules, schema=schema)
        assert search.retrieve(db, '검토확인표식')[0]['id'] == 'CS-DEMO-001/문서'
        new = storage.ingest(db, rows, rules, client, config, schema=schema)
        assert new[0]['effective']['조치내용(2)'] == entry['value']
        changed = [dict(rows[0], 제목='원문 문맥 변경')] + rows[1:]
        new = storage.ingest(db, changed, rules, client, config, schema=schema)
        assert new[0]['reviews']['조치내용(2)']['stale']
        assert all('검토확인표식' not in hit['content'] for hit in search.retrieve(db, '검토확인표식'))
        # 모든 행 처리 실패 시 원문 DB뿐 아니라 검색 인덱스도 함께 보존합니다.
        previous = db.execute('SELECT * FROM rag_documents').fetchall()
        with patch.object(storage, 'normalize', side_effect=ValueError('처리 실패')):
            try:
                storage.ingest(db, rows, rules, client, config, schema=schema)
            except ValueError:
                pass
            else:
                raise AssertionError('처리 실패를 숨겼습니다.')
        assert db.execute('SELECT * FROM rag_documents').fetchall() == previous
    with tempfile.TemporaryDirectory() as tmp:
        output = Path(tmp)
        (output / 'config.json').write_bytes((root / 'config.json').read_bytes())
        with patch.object(cli, 'OpenAI', return_value=client):
            cli.main(output, ['--schema', str(root / 'cslist_schema.json'), '--excel', str(excel), '--question', '배기 튜브 조치는?'])
        result = json.loads((output / 'cslist-output/demo_results.json').read_text(encoding='utf-8'))
        assert result['source_sheet'] == 'CSList' and len(result['cases']) == 1
        page = (output / 'cslist-output/demo.html').read_text(encoding='utf-8')
        assert 'CS-DEMO-001' in page and '보존 열' in page and 'start.bat' in page
        assert not (output / 'demo.sqlite3').exists()
        failing_cases = output / 'failing_cases.json'
        failing_cases.write_text(json.dumps([{'question': 'CS-DEMO-001 관찰 시간은?',
            'expected_ids': ['CS-DEMO-001/작업2'], 'review': '관찰 시간 누락 검사', 'required': ['60\\s*분']}], ensure_ascii=False), encoding='utf-8')
        with patch.object(cli, 'OpenAI', return_value=client):
            try:
                cli.main(output, ['--schema', str(root/'cslist_schema.json'), '--excel', str(excel), '--saved', '--cases', str(failing_cases)])
            except ValueError as error:
                assert '표현 검사' in str(error)
            else:
                raise AssertionError('알려진 답변 누락을 정상 완료로 표시했습니다.')
        # 같은 입력·검토 선택으로 파일 적용까지 실행합니다. 여기서는 API를 쓰지 않습니다.
        (output / 'reviews.json').write_text(json.dumps({'format':'cell-reviews-v1','entries':[entry]},ensure_ascii=False),encoding='utf-8')
        cli.main(output, ['--schema',str(root/'cslist_schema.json'),'--apply-reviews',str(output/'reviews.json')])
        saved = json.loads((output / 'cslist-output/demo_results.json').read_text(encoding='utf-8'))
        assert not saved['cases'] and saved['normalized_records'][0]['effective']['조치내용(2)'] == entry['value']
        # 저장된 원문으로 질문/집계/비교할 때 정리 모델을 다시 호출하지 않습니다.
        with patch.object(cli, 'OpenAI', side_effect=AssertionError('집계·검색 비교가 모델을 호출했습니다.')):
            cli.main(output, ['--schema', str(root/'cslist_schema.json'), '--excel', str(excel), '--saved', '--question', '전체 접수는 총 몇 건인가요?'])
            counted = json.loads((output / 'cslist-output/demo_results.json').read_text(encoding='utf-8'))
            assert counted['cases'][0]['aggregate']['count'] == 8
            cli.main(output, ['--schema', str(root/'cslist_schema.json'), '--excel', str(excel), '--compare-search', '--cases', str(root/'cslist_cases.json')])
            comparison = json.loads((output/'cslist-output/search_comparison.json').read_text(encoding='utf-8'))
            assert len(comparison['cases']) == len(json.loads((root / 'cslist_cases.json').read_text())) and comparison['cases'][0]['changed_cells'] == 1
        # --saved는 등록 시점의 DB 원문만 사용합니다. Excel 변경·삭제는 재등록 전까지 무관합니다.
        with patch.object(cli, 'read_excel', side_effect=AssertionError('DB 조회에서 Excel 재읽기')):
            cli.main(output, ['--schema', str(root/'cslist_schema.json'), '--excel', str(excel), '--saved', '--question', '총 몇 건인가요?'])
        # 메뉴에서 선택한 파일을 유지하며 검토·질문을 같은 CS 결과 폴더에 반영합니다.
        from ragkit import menu
        for name in ('cslist_schema.json', 'cslist_cases.json'):
            (output / name).write_bytes((root / name).read_bytes())
        replies = ['1', str(excel), '2', '6', '2', '3', str(output/'reviews.json'), '4', '배기 튜브 조치는?', '5', '0']
        with patch('builtins.input', side_effect=replies), patch.object(cli, 'OpenAI', return_value=client), patch.object(menu.webbrowser, 'open') as opened:
            cli.main(output, ['--menu'])
        settings = json.loads((output / 'cslist_settings.json').read_text(encoding='utf-8'))
        assert settings['excel'] == str(excel.resolve()) and settings['sheet'] == 'CSList'
        assert opened.call_count == 4
        final = json.loads((output / 'cslist-output/demo_results.json').read_text(encoding='utf-8'))
        assert len(final['cases']) == 1 and final['normalized_records'][0]['effective']['조치내용(2)'] == entry['value']
        # 서버 시간 초과도 실패 사례와 검색 원문을 화면에 남겨야 합니다. 재시도로 대기 시간을 늘리지 않습니다.
        timed_out = Mock(side_effect=APITimeoutError(request=httpx2.Request('POST', 'http://localhost/v1/chat/completions')))
        timeout_client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=timed_out)))
        with patch.object(cli, 'OpenAI', return_value=timeout_client):
            try:
                cli.main(output, ['--schema', str(root/'cslist_schema.json'), '--excel', str(excel),
                                  '--saved', '--question', 'CS-DEMO-001 조치는?'])
            except ValueError:
                pass
            except APITimeoutError as error:
                raise AssertionError('시간 초과가 결과 화면을 남기지 않고 전파됐습니다.') from error
            else:
                raise AssertionError('시간 초과를 성공으로 표시했습니다.')
        failed = json.loads((output/'cslist-output/demo_results.json').read_text(encoding='utf-8'))
        assert failed['cases'][0]['error'] and failed['cases'][0]['hits']
        assert timed_out.call_count == 1
        assert '시간 초과' in (output/'cslist-output/demo.html').read_text(encoding='utf-8')
        # Mac 설정과 출력 경로를 명시해도 Windows 기본 설정·결과를 덮어쓰지 않습니다.
        mac_config = output / 'mac_config.json'
        mac_config.write_text(json.dumps(config | {'base_url': 'http://127.0.0.1:1234/v1', 'model': 'qwen/qwen3.8-27b', 'reasoning_effort': 'none'}), encoding='utf-8')
        windows_before = (output/'config.json').read_bytes()
        with patch.object(cli, 'OpenAI', return_value=client):
            cli.main(output, ['--schema', str(root/'cslist_schema.json'), '--excel', str(excel),
                             '--config', str(mac_config), '--output', str(output/'mac-results'), '--question', 'CS-DEMO-001 조치는?'])
        mac_result = json.loads((output/'mac-results/demo_results.json').read_text(encoding='utf-8'))
        assert mac_result['model'] == 'qwen/qwen3.8-27b' and mac_result['base_url'] == 'http://127.0.0.1:1234/v1'
        assert (output/'config.json').read_bytes() == windows_before
    assert excel.read_bytes() == before
    print('PASS: CS 128열·접수번호·작업별 LLM/검색 분리·FTS5·검토 반영·원문 변경·실패 원자성·독립 출력')
