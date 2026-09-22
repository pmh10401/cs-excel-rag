"""실제 CSList의 열 이름 표기 차이와 중복·오인식 방지를 재현합니다."""
import json
import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from ragkit import records
from ragkit.input_formats import HISTORY_COLUMNS, HISTORY_SHEET
from ragkit.source_cache import convert_excel


def check_blank_columns():
    with tempfile.TemporaryDirectory() as tmp:
        path, cache = Path(tmp)/'headers.xlsx', Path(tmp)/'headers.rag.jsonl.gz'
        book=Workbook(); sheet=book.active; sheet.title=HISTORY_SHEET
        sheet.append(HISTORY_COLUMNS)
        sheet.append(['CS-001','펌프 점검'])
        sheet['AD4'].number_format='0.00'  # 원문·제목 없는 서식뿐인 뒤쪽 열
        book.save(path)
        options={'id_column':'CS NO', 'required_columns':HISTORY_COLUMNS}
        before=path.read_bytes()
        rows=records.read_excel(path,HISTORY_SHEET,**options)
        assert len(rows)==1 and len(records.public_cells(rows[0]))==28
        assert rows[0]['_cell_locations']['설비 타입'].endswith('!AB2')
        convert_excel(path,cache,HISTORY_SHEET,**options)
        assert records.read_excel(cache,HISTORY_SHEET,**options)==rows
        assert path.read_bytes()==before
        cached=cache.read_bytes()
        def reject(expected):
            book.save(path)
            for operation in (lambda:records.read_excel(path,HISTORY_SHEET,**options),
                              lambda:convert_excel(path,cache,HISTORY_SHEET,**options)):
                try: operation()
                except ValueError as error:
                    assert all(part in str(error) for part in expected),str(error)
                else: raise AssertionError('제목 없는 자료나 중복 열을 버리고 읽었습니다.')
            assert cache.read_bytes()==cached
        for value in ('뒤쪽 자료',0,False,'=1+1'):
            sheet['AD3']=value
            reject(['AD1','AD3'])
        sheet['AD3']=None
        for name,expected in [(None,['B1','이름']),('CS NO',['A1','B1','중복']),
                              ('_내부',['B1','예약']), (3,['B1','문자'])]:
            sheet['B1']=name
            reject(expected)
        book.close()

ALIASES = {
    'Part Desc(1)': 'Part Decs(1)', 'Part Desc(2)': 'Part Decs(2)',
    'UNIT(CIS)': 'UNIT(CS)', 'Line': 'line',
    '공수1(1)': '공수(1)', '공수1(2)': '공수(2)',
    '도급여부 CASE': '도급여부CASE', '무상자재개수': '무상자개개수',
    '불출위치(1)': '분출위치(1)', '불출위치(2)': '분출위치(2)',
    '요청 납기일': '요청납기일',
    'CS조치결재상태(조치1/조치2)': '조치결재상태(조치1/조치2)',
    '협력시통합정산CS접수번호': '협력사통합전산CS접수번호',
}


def main(root):
    check_blank_columns()
    schema = json.loads((root / 'cslist_schema.json').read_text(encoding='utf-8'))
    expected = records.read_excel(root / 'cslist_sample.xlsx', 'CSList',
                                  required_columns=schema['columns'], id_column=schema['id_column'])
    names = {old: new for new, old in ALIASES.items()}
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / 'actual-headers.xlsx'
        book = load_workbook(root / 'cslist_sample.xlsx')
        sheet = book['CSList']
        for cell in sheet[1]:
            cell.value = names.get(cell.value, cell.value)
        book.save(path)
        before = path.read_bytes()

        def read(aliases=None):
            return records.read_excel(path, 'CSList', required_columns=schema['columns'],
                                      id_column=schema['id_column'], header_aliases=aliases)

        actual = read()
        assert path.read_bytes() == before, '헤더 대응 과정에서 원본 파일을 수정했습니다.'
        assert len(actual) == len(expected)
        for row, original in zip(actual, expected):
            assert records.public_cells(row) == records.public_cells(original)
            assert row['_typed_cells'] == original['_typed_cells']
            assert records.row_revision(row) == records.row_revision(original)
            for alias, canonical in ALIASES.items():
                assert row['_source_columns'][canonical] == alias
                assert row['_cell_locations'][canonical].split(' / ')[1] == original['_cell_locations'][canonical].split(' / ')[1]
        # 같은 내부 열에 대해 '협력시'와 '협력사' 두 현장 표기를 각각 확인합니다.
        canonical = '협력사통합전산CS접수번호'
        sheet.cell(1, schema['columns'].index(canonical) + 1, '협력사통합정산CS접수번호')
        book.save(path)
        alternate = read()[0]
        assert records.public_cells(alternate) == records.public_cells(expected[0])
        assert alternate['_source_columns'][canonical] == '협력사통합정산CS접수번호'
        assert records.row_revision(alternate) == records.row_revision(expected[0])
        # 비슷한 문자열을 추정해 연결하거나, 두 열을 하나로 덮어쓰면 안 됩니다.
        position = schema['columns'].index('Part Decs(1)') + 1
        sheet.cell(1, position, 'Part Description(1)')
        book.save(path)
        try:
            read()
        except ValueError as error:
            assert 'Part Description(1)' in str(error) and 'Part Decs(1)' in str(error)
        else:
            raise AssertionError('등록되지 않은 이름을 임의로 인식했습니다.')
        # 표기가 다른 식별 열도 사용자가 명시적으로 연결할 수 있어야 합니다.
        sheet.cell(1, schema['columns'].index(schema['id_column']) + 1, '고객 접수 ID')
        book.save(path)
        try:
            read()
        except records.HeaderMappingError as error:
            assert set(error.unknown) == {'고객 접수 ID', 'Part Description(1)'}
            assert schema['id_column'] in error.missing
        aliases = {'고객 접수 ID': schema['id_column'], 'Part Description(1)': 'Part Decs(1)'}
        before = path.read_bytes()
        connected = read(aliases)
        assert records.public_cells(connected[0]) == records.public_cells(expected[0])
        assert connected[0]['_typed_cells'] == expected[0]['_typed_cells']
        assert connected[0]['_source_columns'][schema['id_column']] == '고객 접수 ID'
        assert path.read_bytes() == before
        # 사용자 규칙끼리 같은 열을 가리켜도 입력 시 두 열을 합쳐서는 안 됩니다.
        try:
            read(aliases | {'고객 접수 ID': 'Part Decs(1)'})
        except ValueError as error:
            assert '중복' in str(error)
        else:
            raise AssertionError('사용자 대응의 열 충돌을 허용했습니다.')
        saved = Path(tmp) / 'column_aliases.json'
        records.save_header_aliases(saved, aliases, schema['columns'], schema['id_column'], schema.get('name'))
        assert records.load_header_aliases(saved, schema['columns'], schema['id_column'], schema.get('name')) == aliases
        for bad in ([], {' _reserved': '현상'}, {'현상': '원인(1)'}, {'Line': '원인(1)'}, {'고객 접수 ID': '없는 열'},
                    {'고객 접수 ID': None}, {'고객 접수 ID': '현상', ' 고객 접수 ID ': '원인(1)'}):
            try:
                read(bad)
            except ValueError:
                pass
            else:
                raise AssertionError(f'잘못된 열 연결 허용: {bad}')
        for content in ('{broken', json.dumps({'format': 'old-version'})):
            saved.write_text(content)
            try:
                records.load_header_aliases(saved, schema['columns'], schema['id_column'])
            except ValueError:
                pass
            else:
                raise AssertionError('손상된 열 연결 설정을 무시했습니다.')
        sheet.cell(1, schema['columns'].index(schema['id_column']) + 1, schema['id_column'])
        sheet.cell(1, position, 'Part Desc(1)')
        sheet.cell(1, len(schema['columns']) + 1, 'Part Decs(1)')
        book.save(path)
        try:
            read()
        except ValueError as error:
            assert '중복' in str(error)
        else:
            raise AssertionError('동일 열을 가리키는 두 열을 합쳤습니다.')
        book.close()
    print('PASS: CSList 실제 헤더 표기·기존 양식 호환·원문/자료형/주소 보존·미등록/중복 거절')


if __name__ == '__main__':
    main(Path(__file__).resolve().parents[1] / 'windows-demo')
