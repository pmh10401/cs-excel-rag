"""Excel 연결 경계만 모의 객체로 대체해 xlwings 입력과 공통 검사를 확인합니다."""
import json
import sqlite3
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import Mock, patch
from datetime import date, time, timedelta
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel, CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
from ragkit import records, search, storage, cli


def main(root):
    path = root / 'cslist_sample.xlsx'
    schema = json.loads((root / 'cslist_schema.json').read_text())
    expected = records.read_excel(path, None, required_columns=schema['columns'], id_column='CS 접수번호')
    source = load_workbook(path)
    cells = list(source.active.iter_rows())
    flags = NS(merged=False, formula=False)
    raw_overrides = {}
    def excel_range(start, end=None):
        if end is None:
            r, c = start
            cell = cells[r-1][c-1]
            raw = cell.value
            if isinstance(raw, (date, time, timedelta)):
                raw = to_excel(raw, CALENDAR_MAC_1904 if book.api.Date1904 else CALENDAR_WINDOWS_1900)
            raw = raw_overrides.get(start, raw)
            return NS(api=NS(Value2=raw), number_format=cell.number_format)
        r1, c1 = start; r2, c2 = end
        values = [[cell.value for cell in row[c1-1:c2]] for row in cells[r1-1:r2]]
        block = NS(value=values, api=NS(MergeCells=flags.merged, HasFormula=flags.formula))
        block.options = Mock(return_value=block)
        return block
    sheet = NS(name=source.active.title, range=Mock(side_effect=excel_range),
               api=NS(Cells=NS(Find=Mock(side_effect=lambda **kw: NS(Row=len(cells), Column=len(cells[0]))))))
    book = NS(fullname=str(path.resolve()), sheets=[sheet], api=NS(Date1904=False),
              app=NS(api=NS(CalculationState=0)), save=Mock(), close=Mock())
    app = NS(books=[book], quit=Mock())
    xw = NS(apps=[app], Book=Mock(side_effect=AssertionError('파일 자동 열기 금지')),
            App=Mock(side_effect=AssertionError('Excel 자동 실행 금지')))
    events = []
    with patch.dict(sys.modules, xlwings=xw), patch('sys.platform', 'win32'):
        actual = records.read_excel(path, None, required_columns=schema['columns'], id_column='CS 접수번호', reader='xlwings', progress=events.append)
    assert actual == expected, '읽기 방식만 바꿔도 원문·자료형·검토 버전이 같아야 합니다.'
    assert events[0]['stage'] == '열린 Excel에 연결 중'
    assert events[-1]['processed_rows'] == events[-1]['total_rows'] == len(cells) - 1
    assert any(e['stage'] == 'Excel 값 읽는 중' for e in events)
    from ragkit.source_cache import convert_excel
    with tempfile.TemporaryDirectory() as tmp, patch.dict(sys.modules, xlwings=xw), patch('sys.platform', 'win32'):
        cache = Path(tmp) / 'drm.rag.jsonl.gz'
        events.clear()
        with patch.object(records, 'load_workbook', side_effect=AssertionError('DRM 파일 직접 읽기 금지')):
            convert_excel(path, cache, id_column='CS 접수번호', reader='xlwings', progress=events.append)
            assert records.read_excel(cache, 'CSList', id_column='CS 접수번호') == expected
        assert events[-1]['percent'] == 100 and events[-1]['completed'] == len(expected)
        assert any(e['phase'] == 'reading' for e in events)
        assert all(e['percent'] != 100 for e in events[:-1])
    # 실제 현장 양식의 철자 차이도 파일 읽기와 동일하게 연결합니다.
    from tests.test_headers import ALIASES
    incoming = {canonical: alias for alias, canonical in ALIASES.items()}
    old_headers = [cell.value for cell in cells[0]]
    for cell in cells[0]:
        cell.value = incoming.get(cell.value, cell.value)
    with patch.dict(sys.modules, xlwings=xw), patch('sys.platform', 'win32'):
        alternate = records.read_excel(path, None, required_columns=schema['columns'],
                                       id_column='CS 접수번호', reader='xlwings')
    assert all(records.public_cells(a) == records.public_cells(b) for a, b in zip(alternate, expected))
    assert alternate[0]['_source_columns']['Part Decs(1)'] == 'Part Desc(1)'
    cells[0][schema['columns'].index('CS 접수번호')].value = '현장 접수번호'
    with patch.dict(sys.modules, xlwings=xw), patch('sys.platform', 'win32'):
        custom = records.read_excel(path, None, required_columns=schema['columns'],
                                    id_column='CS 접수번호', reader='xlwings',
                                    header_aliases={'현장 접수번호': 'CS 접수번호'})
    assert records.public_cells(custom[0]) == records.public_cells(expected[0])
    assert custom[0]['_source_columns']['CS 접수번호'] == '현장 접수번호'
    for cell, header in zip(cells[0], old_headers):
        cell.value = header
    from ragkit import excel_app
    with patch.dict(sys.modules, xlwings=xw), patch('sys.platform', 'win32'):
        # Excel의 숫자 변환 때문에 0이 사라지거나 소수·시간이 반올림되지 않아야 합니다.
        headers = [cell.value for cell in cells[0]]
        for name, value, fmt in [('무상자재금액',123.123456,'"₩"0.000000'),
                                  ('초도부품불량',False,'General'),
                                  ('시작시간(1)',time(9,30),'hh:mm:ss'),
                                  ('공수2(1)',timedelta(hours=27,minutes=30),'[h]:mm:ss')]:
            cell = cells[1][headers.index(name)]
            cell.value, cell.number_format = value, fmt
        typed = records.read_excel(path, None, id_column='CS 접수번호', reader='xlwings')[0]
        assert typed['무상자재금액']=='123.123456' and typed['_typed_cells']['초도부품불량']['value'] is False
        assert typed['시작시간(1)']=='09:30:00' and typed['_typed_cells']['공수2(1)']['kind']=='timedelta'
        json.dumps(typed, ensure_ascii=False)  # 누적 시간도 DB에 저장할 수 있어야 합니다.
        book.api.Date1904=True
        assert records.read_excel(path, None, id_column='CS 접수번호', reader='xlwings')[0]['접수일시']==expected[0]['접수일시']
        book.api.Date1904=False
        # 입력 메뉴에서 xlwings 선택과 CSList 시트가 다음 실행에도 유지됩니다.
        from ragkit import menu
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            (folder/'cslist_schema.json').write_bytes((root/'cslist_schema.json').read_bytes())
            with patch('builtins.input',side_effect=['1',str(path.resolve()),'1','6','0']):
                menu.main(folder)
            settings=json.loads((folder/'cslist_settings.json').read_text())
            assert settings['reader']=='xlwings' and settings['sheet']=='CSList'
            before=sheet.range.call_count
            with patch('builtins.input',side_effect=['6','0']):
                menu.main(folder)
            assert sheet.range.call_count>before
        # CS 검색까지 실제 공통 코드로 연결합니다. DRM 파일을 직접 여는 코드는 호출되면 안 됩니다.
        with patch.object(records, 'load_workbook', side_effect=AssertionError('DRM 파일 직접 읽기 금지')):
            with sqlite3.connect(':memory:') as db:
                rules = {key: 'preserve' for key in schema['columns']}
                storage.ingest(db, actual, rules, schema=schema)
                assert 'CS-DEMO-001/작업2' in search.retrieve(db, '배기 튜브 균열 교체')[0]['evidence_ids']
            cli.main(root, ['--schema',str(root/'cslist_schema.json'),'--excel',str(path),'--sheet',sheet.name,
                            '--reader','xlwings','--search-only','--question','배기 튜브 교체'])
        def rejected(message):
            try:
                records.read_excel(path, None, required_columns=schema['columns'], id_column='CS 접수번호', reader='xlwings')
            except (ValueError, RuntimeError) as error:
                assert message in str(error), str(error)
            else:
                raise AssertionError('잘못된 Excel 연결/입력을 허용했습니다: '+message)
        flags.merged = None  # 혼합 병합 범위는 거절하지만 수식 혼합 범위는 계산값을 읽습니다.
        rejected('병합')
        flags.merged = False
        for formula in (True, None):
            flags.formula = formula
            calculated = records.read_excel(path, None, id_column='CS 접수번호', reader='xlwings')
            assert calculated[0]['무상자재금액']=='123.123456'
            assert calculated[0]['시작시간(1)']=='09:30:00'
            assert calculated[0]['제목']==cells[1][1].value
            for state in (1, 2):
                book.app.api.CalculationState = state
                rejected('계산')
            book.app.api.CalculationState = 0
        with tempfile.TemporaryDirectory() as tmp:
            cache = Path(tmp) / 'formula.rag.jsonl.gz'
            convert_excel(path, cache, id_column='CS 접수번호', reader='xlwings')
            assert records.read_excel(cache, 'CSList', id_column='CS 접수번호')==calculated
        flags.formula = False
        old = cells[1][0].value
        cells[1][0].value = cells[2][0].value
        rejected('중복')
        cells[1][0].value = old
        # 오류 셀과 '#N/A'라고 직접 적은 문자는 다릅니다. 오류가 빈 셀로 사라지면 안 됩니다.
        cells[1][1].value = '#N/A'
        raw_overrides[(2,2)] = -2146826246
        rejected('오류')
        raw_overrides.clear()
        assert records.read_excel(path, None, id_column='CS 접수번호', reader='xlwings')[0]['제목'] == '#N/A'
        app.books=[]
        rejected('먼저')
        app.books=[book,book]
        rejected('여러 Excel')
        app.books=[book]
        sheet.range.side_effect=RuntimeError('접근 거절')
        rejected('Excel에서 읽지 못했습니다')
    with patch('sys.platform','darwin'):
        try:
            records.read_excel(path, reader='xlwings')
        except RuntimeError as error:
            assert 'Windows' in str(error)
        else:
            raise AssertionError('Windows 전용 연결을 다른 OS에서 실행했습니다.')
    book.save.assert_not_called(); book.close.assert_not_called(); app.quit.assert_not_called()
    xw.Book.assert_not_called(); xw.App.assert_not_called()
    source.close()
    print('PASS: xlwings 계산값·계산 대기·자료형·변환·CS 검색·중복/병합/오류·원문 저장/종료 방지 (COM 모의 검사)')
