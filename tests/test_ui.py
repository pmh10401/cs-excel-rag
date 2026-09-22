"""임시 Excel·SQLite·모의 모델로 로컬 웹 UI의 실제 HTTP 흐름을 검사합니다."""
import hashlib
import http.client
import json
import sqlite3
import tempfile
import threading
import time
from contextlib import closing
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook


def main(root):
    from ragkit import web
    from ragkit import english

    # SQLite의 with는 트랜잭션만 끝냅니다. closing으로 연결도 닫아 Windows의 임시 DB 삭제를 허용합니다.
    with tempfile.TemporaryDirectory() as folder:
        folder = Path(folder)
        schema_path = root / 'cslist_schema.json'
        schema = json.loads(schema_path.read_text(encoding='utf-8'))
        excel = folder / 'source.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = 'CSList'
        sheet.append(schema['columns'])
        values = {schema['id_column']: 'UI-DEMO-001', '현상': '<script>window.injected=true</script>',
                  '원인(1)': '필터 막힘 의심.', '조치내용(1)': '필터 세척. 유량 11.8 L/min. 필터 교환 안함.'}
        sheet.append([values.get(key, '') for key in schema['columns']])
        second = {schema['id_column']: 'UI-DEMO-002', '제목': '펌프 누수 점검 기록'}
        sheet.append([second.get(key, '') for key in schema['columns']])
        book.save(excel)
        book.close()
        original_hash = hashlib.sha256(excel.read_bytes()).digest()
        config = folder / 'config.json'
        secret = 'PRIVATE-UI-KEY'
        config.write_text(json.dumps({'base_url': 'http://127.0.0.1:1/v1', 'api_key': secret,
                                     'model': 'test-model', 'reasoning_effort': 'high', 'timeout_seconds': 1}))
        app = web.App(root, config, folder / 'output', schema_path)
        app.auto_english = False  # 이 검사는 기존 수동 흐름입니다. 자동 연결은 test_workflow_auto에서 검사합니다.
        # COM의 S_OK(0), S_FALSE(1)는 모두 초기화 횟수를 늘리므로 해제도 필요합니다.
        ole = SimpleNamespace(CoInitializeEx=Mock(), CoUninitialize=Mock())
        with patch.object(web.ctypes, 'windll', SimpleNamespace(ole32=ole), create=True), \
                patch.object(web.sys, 'platform', 'win32'), patch.object(web, 'read_excel', return_value=[]):
            for code in (0, 1):
                ole.CoInitializeEx.return_value = code
                app.read_rows(excel, 'xlwings')
            assert ole.CoUninitialize.call_count == 2
            ole.CoInitializeEx.return_value = -1
            try:
                app.read_rows(excel, 'xlwings')
            except web.ApiError:
                pass
            else:
                raise AssertionError('COM 초기화 실패를 숨겼습니다.')
            assert ole.CoUninitialize.call_count == 2
        server = web.ThreadingHTTPServer(('127.0.0.1', 0), web.make_handler(app))
        app.port = server.server_port
        threading.Thread(target=server.serve_forever, daemon=True).start()
        origin = f'http://127.0.0.1:{app.port}'
        mode = ['valid']
        entered, release = threading.Event(), threading.Event()
        answer_questions = []

        def model_reply(**kwargs):
            if mode[0] == 'slow':
                entered.set()
                assert release.wait(5), '동시 실행 검사 해제가 도착하지 않았습니다.'
            user = kwargs['messages'][1]['content']
            if '원문 자료:\n' in user:
                answer_questions.append(user.split('원문 자료:\n')[0])
            if kwargs['messages'][0]['content'] == english.TRANSLATE_PROMPT:
                data = {key: value + ' maintenance' for key, value in json.loads(user).items()}
            elif kwargs['messages'][0]['content'] == english.QUERY_PROMPT:
                data = {'query': 'maintenance'}
            elif user.startswith('{'):
                payload = json.loads(user)
                data = {key: payload['원본 셀'][key].replace('필터 막힘 의심.', '필터 막힘이 의심됩니다.')
                        for key in payload['수정할 열']}
                if mode[0] == 'numeric_fallback' and '조치내용(1)' in data:
                    data['조치내용(1)'] = '필터 세척. 유량 11.8 L/s. 필터 교환 안함.'
            elif mode[0] == 'insufficient':
                data = {'status': 'insufficient_evidence', 'claims': []}
            else:
                sources = json.loads(user.split('원문 자료:\n')[1])
                source = next(s for s in sources if s['cells'].get('조치내용(1)'))
                quote = source['cells']['조치내용(1)']
                data = {'claims': [{'text': quote if mode[0] != 'invalid' else '999회 확인했습니다.',
                                    'evidence': [{'id': source['id'], 'column': '조치내용(1)'}]}]}
            return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(
                content=json.dumps(data, ensure_ascii=False)))])

        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=model_reply)), close=lambda: None)

        def request(path, payload=None, *, extra=None, status=200, auth=True):
            headers = {'Origin': origin, 'Content-Type': 'application/json'}
            if auth:
                headers.update({'X-CSRF-Token': app.token, 'Cookie': f'csrf_token={app.token}'})
            headers.update(extra or {})
            connection = http.client.HTTPConnection('127.0.0.1', app.port, timeout=10)
            try:
                connection.request('POST' if payload is not None else 'GET', path,
                                   json.dumps(payload) if payload is not None else None, headers)
                response = connection.getresponse()
                content = response.read().decode()
                assert response.status == status, (path, response.status, content)
                assert secret not in content, 'API 응답에 비밀키가 노출되었습니다.'
                return json.loads(content) if response.getheader('Content-Type', '').startswith('application/json') else content
            finally:
                connection.close()

        def finished(expected='done'):
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline:
                job = request('/api/job')
                if job['status'] != 'running':
                    assert job['status'] == expected, job
                    return job
                time.sleep(.02)
            raise AssertionError('백그라운드 작업이 종료되지 않았습니다.')

        try:
            with patch.object(web, 'OpenAI', return_value=client):
                assert 'CS' in request('/')
                assert request('/api/state')['db']['records'] == 0
                request('/api/state', auth=False, status=403)
                request('/api/state', extra={'Host': 'attacker.example'}, status=403)
                request('/api/inspect', {}, extra={'Origin': 'https://attacker.example'}, status=403)
                request('/api/inspect', [], status=400)
                request('/api/overnight', {}, auth=False, status=403)
                request('/api/overnight/stop', {}, auth=False, status=403)
                request('/api/overnight', {'unexpected':True}, status=400)
                request('/api/overnight/stop', {'unexpected':True}, status=400)
                request('/api/overnight/stop', {}, status=409)
                request('/../config.json', status=404)
                request('/api/records?page_size=1000', status=400)
                request('/api/records?page=999999999999999999999', status=400)
                request('/api/inspect', {'path': 42}, status=400)
                request('/api/ask', {'question': ''}, status=400)
                assert request('/api/state')['normalization_workers'] == 200
                request('/api/settings', {'normalization_workers': 2}, auth=False, status=403)
                for invalid in (0, 1001, True, 2.0, '2', None):
                    request('/api/settings', {'normalization_workers': invalid}, status=400)
                request('/api/settings', {'normalization_workers': 2, 'api_key': 'changed'}, status=400)
                for invalid in (-1, 101, True, '3', None, 1.5):
                    request('/api/settings', {'reference_limit': invalid}, status=400)
                    request('/api/ask', {'question': '현상', 'reference_limit': invalid}, status=400)
                for limit in (0, 1, 100, 3):
                    assert request('/api/settings', {'reference_limit': limit})['reference_limit'] == limit
                    assert request('/api/state')['settings']['reference_limit'] == limit
                    restored = web.App(root, config, folder / 'output', schema_path)
                    assert restored.reference_limit == 3  # 질문 설정은 사용자 세션별이며 브라우저에 보관합니다.
                assert request('/api/settings', {'normalization_workers': 1000})['normalization_workers'] == 1000
                assert web.App(root, config, folder / 'output', schema_path).normalization_workers == 1000
                assert request('/api/settings', {'normalization_workers': 2})['normalization_workers'] == 2
                assert web.App(root, config, folder / 'output', schema_path).normalization_workers == 2
                with patch.object(app, 'save_settings', side_effect=OSError('저장 실패')):
                    request('/api/settings', {'normalization_workers': 3}, status=500)
                assert app.normalization_workers == 2
                assert web.App(root, config, folder / 'output', schema_path).normalization_workers == 2
                # 기존 입력 검사·저장도 동시 처리 수를 지우지 않습니다.
                selection = {'path': str(excel), 'reader': 'openpyxl'}
                invalid_headers = folder / 'duplicate-headers.xlsx'
                book = load_workbook(excel)
                book['CSList']['B1'] = book['CSList']['A1'].value
                book.save(invalid_headers)
                book.close()
                detail = request('/api/inspect', selection | {'path': str(invalid_headers)}, status=400)
                assert all(part in detail['error'] for part in ('duplicate-headers.xlsx', 'CSList', 'A1, B1', '중복'))
                assert not app.db_path.exists(), '입력 오류 검사에서 DB를 생성했습니다.'
                alternate = folder / 'custom-headers.xlsx'
                book = load_workbook(excel)
                book['CSList'].cell(1, schema['columns'].index('현상') + 1, '현장 증상 기록')
                book.save(alternate)
                book.close()
                custom_selection = selection | {'path': str(alternate)}
                detail = request('/api/inspect', custom_selection, status=400)
                assert detail['code'] == 'header_mapping_required'
                assert detail['mapping']['missing'] == ['현상']
                assert detail['mapping']['unknown'] == ['현장 증상 기록']
                aliases = {'현장 증상 기록': '현상'}
                request('/api/header-mappings', {'aliases': aliases}, auth=False, status=403)
                request('/api/header-mappings', {'aliases': {'현상': '원인(1)'}}, status=400)
                assert request('/api/header-mappings', {'aliases': aliases})['saved']
                assert request('/api/inspect', custom_selection)['rows'] == 2
                reopened = web.App(root, config, folder / 'output', schema_path)
                assert reopened.header_aliases == aliases
                assert request('/api/state')['header_aliases'] == aliases
                # 파일을 읽지 못해도 잘못 저장한 대응은 삭제할 수 있습니다.
                alternate.unlink()
                assert request('/api/header-mappings', {'aliases': {}})['aliases'] == {}
                assert request('/api/inspect', selection)['rows'] == 2
                assert web.App(root, config, folder / 'output', schema_path).normalization_workers == 2
                assert not app.db_path.exists(), '입력 검사만으로 DB를 생성했습니다.'
                request('/api/convert', selection, auth=False, status=403)
                # Excel이 응답을 기다리는 동안에도 상태 API가 단계·경과 시간을 반환합니다.
                convert_entered, convert_release = threading.Event(), threading.Event()
                convert = web.convert_excel
                def paused_convert(*args, progress, **options):
                    def report(info):
                        progress(info)
                        if info['phase'] == 'opening':
                            convert_entered.set()
                            assert convert_release.wait(10)
                    return convert(*args, progress=report, **options)
                with patch.object(web, 'convert_excel', side_effect=paused_convert):
                    try:
                        request('/api/convert', selection, status=202)
                        assert convert_entered.wait(5)
                        waiting = request('/api/job')
                        assert waiting['status'] == 'running' and waiting['progress']['phase'] == 'opening'
                        assert waiting['progress']['percent'] is None
                        time.sleep(.12)
                        assert request('/api/job')['elapsed_seconds'] > waiting['elapsed_seconds']
                    finally:
                        convert_release.set()
                    completed_conversion = finished()
                assert completed_conversion['progress']['percent'] == 100
                assert completed_conversion['progress']['completed'] == 2
                converted = completed_conversion['result']
                assert converted['rows'] == 2 and Path(converted['path']).is_file()
                cached_selection = {'path': converted['path'], 'reader': 'xlwings'}
                assert request('/api/state')['excel'] == converted['path']
                assert web.App(root, config, folder / 'output', schema_path).excel_text == converted['path']
                assert not app.db_path.exists(), '변환만으로 DB를 생성했습니다.'
                request('/api/convert', cached_selection, status=400)
                with patch('ragkit.records.load_workbook', side_effect=AssertionError('Excel 재조회')):
                    assert request('/api/inspect', cached_selection)['rows'] == 2
                mode[0] = 'numeric_fallback'
                request('/api/normalize', cached_selection, status=202)
                completed = finished()
                mode[0] = 'valid'
                assert completed['result']['rows'] == 2
                assert completed['result']['fallback_cells'] == completed['progress']['fallback_cells'] == 1
                assert completed['progress']['completed'] == completed['progress']['total'] == 2
                time.sleep(.12)
                assert app.job_snapshot()['elapsed_seconds'] == completed['elapsed_seconds'], '완료 시간이 계속 늘어납니다.'
                listing = request('/api/records?page_size=1')
                assert listing['total'] == 2 and len(listing['items']) == 1
                assert request('/api/records?page_size=1&page=2')['items'][0]['id'] == 'UI-DEMO-002'
                record = request('/api/record?id=UI-DEMO-001')
                cells = {c['column']: c for g in record['groups'] for c in g['cells']}
                assert cells['원인(1)']['proposal'] == '필터 막힘이 의심됩니다.'
                assert cells['원인(1)']['action'] == 'pending'
                assert cells['조치내용(1)']['fallback']
                assert cells['조치내용(1)']['proposal'] == cells['조치내용(1)']['original']
                assert not cells['원인(1)']['fallback']
                request('/api/llm-messages', auth=False, status=403)
                exchanges = request('/api/llm-messages')['items']
                assert exchanges and any(item['kind'] == 'normalize' for item in exchanges)
                assert 'request' not in exchanges[0]
                traced = [request('/api/llm-messages?id=' + item['id']) for item in exchanges]
                assert any('보정 참고' in item['request'] for item in traced)
                assert all(secret not in json.dumps(item) for item in traced)
                assert any('필터' in item['response'] for item in traced)
                request('/api/llm-messages?id=unknown', status=404)
                assert not cells['MTTR']['editable'] and cells['MTTR']['original'] == ''
                entry = {'source_key': record['source_key'], 'id': record['id'], 'revision': record['revision'],
                         'column': '원인(1)', 'action': 'accept', 'value': cells['원인(1)']['proposal']}
                review = {'format': 'cell-reviews-v1', 'entries': [entry]}
                second = entry | {'column': '조치내용(1)', 'action': 'original', 'value': cells['조치내용(1)']['original']}
                # 일괄 요청 중 하나라도 검사가 실패하면 나머지도 저장하지 않습니다.
                request('/api/reviews', {'format': 'cell-reviews-v1', 'entries': [entry, second | {'revision': 'old'}]}, status=400)
                unchanged = request('/api/record?id=UI-DEMO-001')
                assert all(c['action'] == 'pending' for g in unchanged['groups'] for c in g['cells'])
                assert request('/api/reviews', {'format': 'cell-reviews-v1', 'entries': [entry, second]})['saved'] == 2
                request('/api/reviews', {'format': 'cell-reviews-v1', 'entries': [second | {'action': 'pending'}]})
                assert request('/api/reviews', review)['saved'] == 1
                assert request('/api/inspect', selection)['matches_stored'] is True
                assert request('/api/inspect', cached_selection)['matches_stored'] is True
                request('/api/inspect', selection)
                record = request('/api/record?id=UI-DEMO-001')
                cells = {c['column']: c for g in record['groups'] for c in g['cells']}
                assert cells['원인(1)']['proposal'] == entry['value'] and cells['원인(1)']['action'] == 'accept'
                request('/api/ask', {'question': 'UI-DEMO-001의 조치내용은?'}, status=202)
                result = finished()['result']
                evidence = result['claims'][0]['evidence'][0]
                assert result['status'] == 'ok' and evidence['quote'] == values['조치내용(1)']
                assert 'CSList!' in evidence['location'] and '11.8 L/min' in result['answer']
                assert '원문 보완' not in result['answer']
                assert values['조치내용(1)'] in result['answer']
                assert result['claims'][-1]['evidence'][-1]['quote'] == values['조치내용(1)']
                assert len(result['claims']) == 1
                cited = {(e['id'], e['column']) for c in result['claims'] for e in c['evidence']}
                assert all((e['id'], e['column']) not in cited for c in result['source_supplements'] for e in c['evidence'])
                mode[0] = 'insufficient'
                request('/api/ask', {'question': 'UI-DEMO-001의 조치내용은?'}, status=202)
                declined = finished()['result']
                assert declined['status'] == 'insufficient_evidence' and declined['error'] is None
                assert declined['claims'] == [] and '검색된 후보' in declined['answer']
                mode[0] = 'valid'
                from ragkit import cslist
                with patch.object(app, 'read_rows', side_effect=AssertionError('변경 없는 파일 재읽기')), \
                        patch.object(cslist, 'rebuild_search', side_effect=AssertionError('질문마다 색인 재생성')):
                    request('/api/ask', {'question': 'UI-DEMO-001의 조치내용은?'}, status=202)
                    cached = finished()['result']
                assert cached['source_db'] and not cached['index_rebuilt']
                assert {'source_check', 'index_prepare', 'keyword', 'answer', 'total'} <= cached['timings'].keys()
                request('/api/ask', {'question': 'UI-DEMO-999의 조치내용은?'}, status=202)
                assert finished()['result']['status'] == 'no_match'
                with patch.object(web, 'OpenAI', side_effect=AssertionError('없는 증상에 모델을 호출했습니다.')):
                    request('/api/ask', {'question': '펌프 소음 과거 조치는?'}, status=202)
                    assert finished()['result']['status'] == 'no_match'
                mode[0] = 'invalid'
                request('/api/ask', {'question': 'UI-DEMO-001의 조치내용은?'}, status=202)
                failed_answer = finished('error')
                assert failed_answer['code'] == 'validation_failed'
                events = failed_answer['result']['validation_events']
                assert len(events) == 2 and events[-1]['outcome'] == 'failed'
                assert all('response' not in event for event in events)
                assert failed_answer['result']['claims'] == []
                assert (app.output/'answer_diagnostics.jsonl').exists()
                mode[0] = 'slow'
                request('/api/ask', {'question': 'UI-DEMO-001의 조치내용은?'}, status=202)
                assert entered.wait(5)
                request('/api/ask', {'question': '두 번째 질문'}, status=409)
                request('/api/reviews', review, status=409)
                request('/api/reviews/accept-pending', {}, status=409)
                request('/api/english', {}, status=409)
                request('/api/header-mappings', {'aliases': {}}, status=409)
                request('/api/settings', {'normalization_workers': 1}, status=409)
                release.set()
                finished()
                mode[0] = 'valid'
                assert hashlib.sha256(excel.read_bytes()).digest() == original_hash, 'UI가 원본 Excel을 수정했습니다.'
                reopened = web.App(root, config, folder / 'output', schema_path)
                assert Path(reopened.excel_text).resolve() == excel.resolve()
                with closing(sqlite3.connect(app.db_path)) as db, db:
                    assert db.execute('SELECT action,value FROM cell_reviews').fetchone() == ('accept', entry['value'])
                book = load_workbook(excel)
                book['CSList'].cell(2, schema['columns'].index('원인(1)') + 1, '필터 막힘 확인.')
                book.save(excel)
                book.close()
                request('/api/ask', {'question': 'UI-DEMO-001의 조치내용은?'}, status=202)
                assert finished()['result']['source_db'], '파일 수정이 등록된 DB 질문을 차단했습니다.'
                with closing(sqlite3.connect(app.db_path)) as db, db:
                    assert db.execute('SELECT value FROM cell_reviews').fetchone()[0] == entry['value']
                # UI에서도 부분 완료 수와 실패 ID를 보존하고 같은 요청으로 재개합니다.
                from ragkit import storage, records
                app.config['model'] = 'different-model'
                request('/api/settings', {'normalization_workers': 1})
                calls = []
                def partial(_client, _config, row, _rules, schema=None):
                    calls.append(records.record_id(row))
                    if len(calls) == 1 and mode[0] != 'retry':
                        raise ValueError('의도한 행 실패 ' + secret)
                    return records.public_cells(row)
                with patch.object(storage, 'normalize', side_effect=partial):
                    request('/api/normalize', selection, status=202)
                    failed = finished('error')
                assert failed['progress']['completed'] == 1
                assert failed['progress']['failed_id'] == 'UI-DEMO-001'
                assert failed['code'] == 'normalization_partial'
                assert failed['progress']['attempted'] == 2 and failed['progress']['failed'] == 1
                assert failed['progress']['failures'] == [{'id': 'UI-DEMO-001', 'error': '의도한 행 실패 ***'}]
                assert calls == ['UI-DEMO-001', 'UI-DEMO-002']
                calls.clear()
                mode[0] = 'retry'
                with patch.object(storage, 'normalize', side_effect=partial):
                    request('/api/normalize', selection, status=202)
                    resumed = finished()
                assert calls == ['UI-DEMO-001']
                assert resumed['progress']['reused'] == 1 and resumed['result']['rows'] == 2
                # DB 전체 채택도 접근 검사·저장 실패 롤백·기존 검토 보존을 거쳐야 합니다.
                request('/api/reviews/accept-pending', {}, auth=False, status=403)
                request('/api/reviews/accept-pending', {}, extra={'Origin': 'https://attacker.example'}, status=403)
                request('/api/reviews/accept-pending', {'entries': []}, status=400)
                with closing(sqlite3.connect(app.db_path)) as db:
                    previous = db.execute('SELECT * FROM cell_reviews').fetchall()
                with patch.object(web, 'accept_pending_reviews', side_effect=RuntimeError('의도한 저장 실패')):
                    request('/api/reviews/accept-pending', {}, status=202)
                    finished('error')
                with closing(sqlite3.connect(app.db_path)) as db:
                    assert db.execute('SELECT * FROM cell_reviews').fetchall() == previous
                entered.clear(); release.clear()
                accept = web.accept_pending_reviews
                def slow_accept(*args, **kwargs):
                    entered.set()
                    assert release.wait(5)
                    return accept(*args, **kwargs)
                with patch.object(web, 'accept_pending_reviews', side_effect=slow_accept), \
                        patch.object(web, 'OpenAI', side_effect=AssertionError('일괄 채택은 LLM을 호출하면 안 됩니다.')):
                    request('/api/reviews/accept-pending', {}, status=202)
                    assert entered.wait(5)
                    request('/api/reviews/accept-pending', {}, status=409)
                    request('/api/reviews', review, status=409)
                    request('/api/normalize', selection, status=409)
                    release.set()
                    result = finished()['result']
                assert result['saved'] == 2 and result['applied_rows'] == 1, result
                assert result['reviewed_cells'] == 1, '이전 버전의 직접 검토도 자동 덮어쓰지 않습니다.'
                updated = request('/api/record?id=UI-DEMO-001')
                cells = {c['column']: c for g in updated['groups'] for c in g['cells']}
                assert cells['조치내용(1)']['action'] == 'accept'
                assert cells['원인(1)']['stale'], '재검토 대상을 일괄 채택으로 숨겼습니다.'
                request('/api/reviews/accept-pending', {}, status=202)
                assert finished()['result']['saved'] == 0
                request('/api/english', {}, auth=False, status=403)
                request('/api/english', {}, extra={'Origin': 'https://attacker.example'}, status=403)
                for limit in (-1, True, '20'):
                    request('/api/english', {'limit': limit}, status=400)
                request('/api/ask', {'question': '현상', 'bilingual': 'yes'}, status=400)
                request('/api/english', {'limit': 1}, status=202)
                result = finished()['result']
                assert result['generated'] == 1 and result['indexed'] == 1, result
                assert request('/api/state')['db']['english_records'] == 1
                request('/api/english', {'limit': 0}, status=202)
                assert finished()['result']['indexed'] == 2
                request('/api/english', {'limit': 0}, status=202)
                assert finished()['result']['processed'] == 0
                mode[0] = 'valid'
                request('/api/ask', {'question': '소모품 정비', 'bilingual': True}, status=202)
                result = finished()['result']
                assert result['search']['mode'] == 'bilingual' and result['search']['english_query'] == 'maintenance'
                assert result['search']['translation_status'] == 'translated'
                assert '소모품 정비' in answer_questions[-1] and 'maintenance' in answer_questions[-1]
                assert result['claims'][0]['evidence'][0]['quote'] == values['조치내용(1)']
                with patch.object(english, 'retrieve', wraps=english.retrieve) as lookup:
                    mode[0] = 'insufficient'  # 1위 후보는 조치가 없는 접수입니다.
                    request('/api/ask', {'question': '소모품 정비', 'bilingual': True, 'reference_limit': 1}, status=202)
                    limited = finished()['result']
                    assert limited['references']['selected'] == 1
                    assert all(call.kwargs['limit'] == 1 for call in lookup.call_args_list)
                mode[0] = 'valid'
                request('/api/ask', {'question': 'UI-DEMO-001의 조치내용은?', 'reference_limit': 0}, status=202)
                automatic = finished()['result']
                assert automatic['references']['mode'] == 'auto'
                assert automatic['references']['context_source'] == 'unknown'
                assert automatic['references']['selected'] == 1
                with patch.dict(app.config, {'max_context_chars': 10}):
                    request('/api/ask', {'question': 'UI-DEMO-001의 조치내용은?', 'reference_limit': 0}, status=202)
                    blocked = finished('error')
                    assert blocked['result']['references']['selected'] == 0
                    assert not blocked['result']['claims']
                request('/api/ask', {'question': '전체 접수는 총 몇 건인가요?', 'reference_limit': 1}, status=202)
                aggregate_result = finished()['result']
                assert aggregate_result['aggregate']['count'] == 2
                assert aggregate_result['references']['applied_to'] == 'aggregate'
                from ragkit import semantic
                app.config['embedding'] = {'base_url': 'http://127.0.0.1:1/v1', 'model': 'mock', 'api_key': 'PRIVATE-EMBEDDING-KEY', 'auto_index': False}
                vector_client = SimpleNamespace(close=lambda: None, embeddings=SimpleNamespace(create=lambda **kw:
                    SimpleNamespace(data=[SimpleNamespace(index=i, embedding=[1., 0.]) for i in range(len(kw['input']))])))
                request('/api/vectors', {}, auth=False, status=403)
                request('/api/compare-search', {}, auth=False, status=403)
                request('/api/vectors', {}, extra={'Origin': 'https://attacker.example'}, status=403)
                for limit in (-1, True, '20'):
                    request('/api/vectors', {'limit': limit}, status=400)
                request('/api/ask', {'question': '증상', 'search_mode': 'invalid'}, status=400)
                with closing(sqlite3.connect(app.db_path)) as db:
                    saved = db.execute('SELECT * FROM normalized_records').fetchall()
                    reviews = db.execute('SELECT * FROM cell_reviews').fetchall()
                with patch.object(semantic, 'client_for', return_value=vector_client):
                    request('/api/vectors', {'limit': 0}, status=202)
                    assert finished()['result']['indexed'] == 2
                    state = request('/api/state')
                    assert state['db']['vector_records'] == 2 and 'PRIVATE-EMBEDDING-KEY' not in json.dumps(state)
                    request('/api/vectors', {'limit': 0}, status=202)
                    assert finished()['result']['generated'] == 0
                    with patch.object(web, 'OpenAI', side_effect=AssertionError('검색 비교에서 답변 모델 호출')):
                        with patch.object(semantic, 'retrieve', wraps=semantic.retrieve) as retrieve, \
                                patch.object(web, 'retrieve', side_effect=AssertionError('키워드 검색 중복 실행')):
                            request('/api/compare-search', {'question': '펌프 물기'}, status=202)
                            result = finished()['result']
                        assert retrieve.call_count == 1
                    assert result['vector_search']['vector_ids'] and result['vector_search']['hybrid_ids']
                with closing(sqlite3.connect(app.db_path)) as db:
                    assert db.execute('SELECT * FROM normalized_records').fetchall() == saved
                    assert db.execute('SELECT * FROM cell_reviews').fetchall() == reviews
                app.config['embedding']['auto_index'] = True
                with patch.object(semantic, 'client_for', return_value=vector_client):
                    app.start_embedding()
                    assert finished()['result']['generated'] == 0  # 시작할 때 이미 만든 벡터 재사용
                    rec = request('/api/record?id=UI-DEMO-001')
                    entry = {'id': rec['id'], 'source_key': rec['source_key'], 'revision': rec['revision'],
                             'column': '현상', 'action': 'edit', 'value': '필터 막힘 의심. 유량 저하.'}
                    assert request('/api/reviews', {'format':'cell-reviews-v1','entries':[entry]})['saved'] == 1
                    assert finished()['result']['embedding']['generated'] == 1
                    assert request('/api/state')['embedding_status']['status'] == 'ready'
                    with patch.object(app, 'update_vectors', side_effect=AssertionError('다듬기 직후 벡터 호출')):
                        request('/api/normalize', {}, status=202)
                        assert 'embedding' not in finished()['result']
                    request('/api/reviews/accept-pending', {}, status=202)
                    bulk_result = finished()['result']
                    if bulk_result['applied_rows']:
                        assert bulk_result['embedding']['indexed'] == 2
                    else:
                        assert 'embedding' not in bulk_result
                    # 변환 뒤 자동 색인은 기존 제안과 검토를 덮어쓰지 않습니다.
                    with closing(sqlite3.connect(app.db_path)) as db:
                        saved_reviews = db.execute('SELECT * FROM cell_reviews').fetchall()
                        after_review = db.execute('SELECT * FROM normalized_records').fetchall()
                    request('/api/convert', selection, status=202)
                    assert finished()['result']['embedding']['generated'] == 0
                    with closing(sqlite3.connect(app.db_path)) as db:
                        assert db.execute('SELECT * FROM normalized_records').fetchall() == after_review
                        assert db.execute('SELECT * FROM cell_reviews').fetchall() == saved_reviews
                with patch.object(semantic, 'client_for', side_effect=RuntimeError('임베딩 연결 실패')):
                    entry['value'] = '필터 막힘 의심. 유량 저하 재발.'
                    assert request('/api/reviews', {'format':'cell-reviews-v1','entries':[entry]})['saved'] == 1
                    job = finished()
                    assert job['kind'] == 'auto_search' and job['result']['embedding_warning']
                    assert request('/api/state')['embedding_status']['status'] == 'error'
                rec = request('/api/record?id=UI-DEMO-001')
                assert next(c for g in rec['groups'] for c in g['cells'] if c['column']=='현상')['value'] == entry['value']
                with patch.object(semantic, 'client_for', return_value=vector_client):
                    app.start_embedding()
                    assert finished()['result']['generated'] == 1
                from ragkit import overnight
                def waiting(*args):
                    assert app.overnight_stop.wait(5)
                    return {'ready':False,'stopped':True}
                with patch.object(overnight,'run',side_effect=waiting):
                    request('/api/overnight',selection,status=202)
                    request('/api/overnight',selection,status=409)
                    request('/api/settings',{'normalization_workers':1},status=409)
                    request('/api/overnight/stop',{},extra={'Origin':'https://attacker.example'},status=403)
                    assert request('/api/overnight/stop',{})['stopping']
                    assert finished()['result']['stopped']
        finally:
            release.set()
            server.shutdown()
            server.server_close()
    print('PASS: 웹 UI HTTP·원문 보존·검토 저장·질문·근거 검사·중복 요청·접근 제한')


if __name__ == '__main__':
    main(Path(__file__).resolve().parents[1] / 'windows-demo')
