"""Excel 입력 경계를 검사합니다. 실행: .venv/bin/python test_excel.py"""
from unittest.mock import patch

from openpyxl import load_workbook

from ragkit import records
from ragkit.local_cli import read_excel


def rejected(excel, change, message):
    # 파일을 저장하지 않고 메모리에 읽은 복사본만 바꾸므로 원본 Excel은 그대로입니다.
    book = load_workbook(excel)
    change(book['점검기록'])
    with patch.object(records, 'load_workbook', return_value=book):
        try:
            read_excel(excel)
        except ValueError as error:
            assert message in str(error), str(error)
        else:
            raise AssertionError(f'잘못된 입력을 허용했습니다: {message}')


def main(root):
    excel = root / 'outputs' / '01a083b4' / 'sample_data.xlsx'
    rows = read_excel(excel)
    assert len(rows) == 6
    assert rows[0]['id'] == 'DEMO-001'
    assert rows[0]['_excel_location'].endswith('점검기록!A2:H2')
    assert rows[1]['기록된 원인'] == '필터 막힘'

    # 첫 두 열을 통째로 바꿔도 위치가 아니라 열 이름으로 정확하게 매핑해야 합니다.
    book = load_workbook(excel)
    sheet = book['점검기록']
    for line in sheet.iter_rows():
        line[0].value, line[1].value = line[1].value, line[0].value
    sheet['G2'] = None
    sheet.insert_rows(3)  # 빈 행은 건너뛰고 실제 Excel 행 번호는 유지해야 합니다.
    with patch.object(records, 'load_workbook', return_value=book):
        changed = read_excel(excel)
    assert len(changed) == 6 and changed[0]['id'] == 'DEMO-001'
    assert changed[0]['비고'] == ''
    assert changed[1]['_excel_location'].endswith('A4:H4')

    rejected(excel, lambda s: setattr(s['A1'], 'value', '문서명'), '첫 행')
    rejected(excel, lambda s: setattr(s['H1'], 'value', None), '첫 행')
    rejected(excel, lambda s: setattr(s['A3'], 'value', 'DEMO-001'), '중복 ID')
    rejected(excel, lambda s: setattr(s['A2'], 'value', None), 'id')
    rejected(excel, lambda s: setattr(s['F2'], 'value', '=1+1'), '수식')
    rejected(excel, lambda s: setattr(s['F2'], 'value', 0), '문자 값')
    rejected(excel, lambda s: s.merge_cells('B2:B3'), '병합')
    rejected(excel, lambda s: s.delete_rows(2, 6), '데이터 행')
    try:
        read_excel(excel, '없는시트')
    except ValueError as error:
        assert '시트를 찾을 수 없습니다' in str(error)
    else:
        raise AssertionError('잘못된 시트 이름을 허용했습니다.')
    print('PASS: 실제 XLSX 읽기·열 순서·빈 셀/행·출처·잘못된 입력 거절')
