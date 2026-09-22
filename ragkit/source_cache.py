"""Excel을 한 번 읽어 반복 실험용 압축 JSONL로 저장합니다. 모델은 호출하지 않습니다."""
import gzip
import json
import math
import posixpath
import re
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter
from .records import iter_rows_from_cells, read_excel, public_cells, record_id, text_for
from .input_formats import select_sheet, input_layout

SUFFIX = '.rag.jsonl.gz'
FORMAT = 'rag-source-v1'


def _check_merges(path, sheet_name, progress):
    # 읽기 전용 openpyxl은 병합 정보를 제공하지 않아 선택한 시트의 XML을 검사합니다.
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with ZipFile(path) as archive:
        sheets = ET.fromstring(archive.read('xl/workbook.xml')).find('s:sheets', ns)
        identity = next(s.attrib['{' + rel_ns + '}id'] for s in sheets if s.attrib['name'] == sheet_name)
        rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        target = next(r.attrib['Target'] for r in rels if r.attrib['Id'] == identity)
        member = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        last_row = 0
        with archive.open(member) as stream:
            for _, element in ET.iterparse(stream, events=('end',)):
                tag = element.tag.rsplit('}', 1)[-1]
                if tag == 'mergeCell':
                    raise ValueError('입력 시트의 병합 셀을 해제하고 각 행에 값을 입력하세요.')
                if tag == 'row':
                    last_row = max(last_row, int(element.get('r', last_row + 1)))
                    if last_row % 1000 == 0:
                        progress({'processed_rows': max(0, last_row - 1)})
                element.clear()
        return max(0, last_row - 1)


def _excel_rows(path, sheet_name, progress, **options):
    if options.get('reader') == 'xlwings':
        yield from read_excel(path, sheet_name, progress=progress, **options)
        return
    if options.pop('reader', 'openpyxl') != 'openpyxl':
        raise ValueError('읽기 방식은 openpyxl 또는 xlwings여야 합니다.')
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet_name = select_sheet(book.sheetnames, sheet_name)
        progress({'phase': 'checking', 'stage': '시트 구조·전체 행 수 확인 중'})
        total = _check_merges(path, sheet_name, progress)
        sheet = book[sheet_name]
        sheet.reset_dimensions()  # 잘못 저장된 사용 범위 때문에 뒤쪽 행을 놓치지 않습니다.
        progress({'phase': 'converting', 'stage': '변환 파일 작성 중', 'total_rows': total, 'processed_rows': 0})
        def cells():
            for number, row in enumerate(sheet.iter_rows(), 1):
                yield row
                # 빈 행도 읽기 진행에 포함하되 변환 건수에는 포함하지 않습니다.
                progress({'processed_rows': max(0, number - 1)})
        yield from iter_rows_from_cells(path, sheet_name, cells(),
            options.get('required_columns'), options.get('allow_typed', True),
            options.get('id_column', 'id'), options.get('header_aliases'))
    finally:
        book.close()


def convert_excel(path, output, sheet_name='CSList', *, progress=None, **options):
    path, output = Path(path), Path(output)
    if path.suffix.lower() != '.xlsx' or not output.name.lower().endswith(SUFFIX):
        raise ValueError('입력은 .xlsx, 변환 파일은 .rag.jsonl.gz로 지정하세요.')
    started = time.monotonic()
    count = 0
    state = {'phase': 'opening', 'stage': 'Excel 입력 준비 중', 'processed_rows': 0, 'total_rows': None}
    last_sent, last_stage = 0, None
    def report(info):
        nonlocal last_sent, last_stage
        state.update(info)
        now = time.monotonic()
        total = state['total_rows']
        percent = min(99, int(100 * state['processed_rows'] / total)) if total else None
        state.update(completed=count, elapsed_seconds=round(now - started, 1),
                     percent=100 if state['phase'] == 'done' else percent)
        # 행마다 상태는 갱신하되 화면/콘솔 전달은 최대 초당 네 번으로 제한합니다.
        if progress and (state['stage'] != last_stage or now - last_sent >= .25):
            progress(dict(state))
            last_sent, last_stage = now, state['stage']
    report({})
    output.parent.mkdir(parents=True, exist_ok=True)
    # 실패한 변환이 기존 파일을 덮지 않도록 같은 폴더에서 완성한 뒤 교체합니다.
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.tmp', delete=False) as handle:
        temporary = Path(handle.name)
    rows = _excel_rows(path, sheet_name, report, **options)
    try:
        with gzip.open(temporary, 'wt', encoding='utf-8', compresslevel=1, newline='\n') as stream:
            def write(value):
                stream.write(json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False) + '\n')
            for row in rows:
                columns = list(public_cells(row))
                if count == 0:
                    sheet_name = row['_source_sheet']
                    write({'format': FORMAT, 'created_at': datetime.now(timezone.utc).isoformat(),
                           'source_file': path.name, 'source_key': row['_source_key'],
                           'sheet': row['_source_sheet'], 'id_column': row['_id_column'],
                           'columns': columns, 'source_columns': row.get('_source_columns', {})})
                number = int(row['_excel_location'].rsplit('!', 1)[1].split(':')[0][1:])
                write({'row': number, 'values': [row[c] for c in columns], 'typed': row['_typed_cells']})
                count += 1
                report({'phase': 'converting', 'stage': '변환 파일 작성 중', 'processed_rows': number - 1})
            write({'end': count})
            report({'phase': 'saving', 'stage': '압축 파일 저장 중'})
        temporary.replace(output)
    finally:
        rows.close()
        temporary.unlink(missing_ok=True)
    report({'phase': 'done', 'stage': '변환 완료'})
    return {'path': str(output.resolve()), 'source': str(path.resolve()), 'rows': count,
            'bytes': output.stat().st_size, 'format': FORMAT, 'sheet': sheet_name}


def iter_cache(path, sheet_name='CSList', *, required_columns=None, allow_typed=True, id_column='id'):
    """파일 끝까지 검사합니다. 목록이 필요 없는 검증·집계는 메모리를 적게 사용합니다."""
    def require(condition):
        if not condition:
            raise ValueError('변환 파일의 형식·양식 또는 원본 정보가 올바르지 않습니다. Excel에서 다시 변환하세요.')
    try:
        with gzip.open(path, 'rt', encoding='utf-8') as stream:
            meta = json.loads(next(stream, '{}'))
            require(isinstance(meta, dict) and meta.get('format') == FORMAT)
            columns, sheet = meta.get('columns'), meta.get('sheet')
            require(isinstance(sheet, str) and bool(sheet))
            sheet_name = select_sheet([sheet], sheet_name)
            required_columns, id_column = input_layout(sheet, required_columns, id_column)
            require(isinstance(columns, list) and bool(columns) and all(isinstance(c, str) and c and not c.startswith('_') for c in columns))
            require(len(set(columns)) == len(columns) and id_column in columns and meta.get('id_column') == id_column)
            require(required_columns is None or set(columns) == set(required_columns))
            require(isinstance(sheet, str) and bool(sheet) and (sheet_name is None or sheet == sheet_name))
            filename, key, aliases = meta.get('source_file'), meta.get('source_key'), meta.get('source_columns')
            require(isinstance(filename, str) and bool(filename) and '/' not in filename and '\\' not in filename)
            require(isinstance(key, str) and re.fullmatch('[0-9a-f]{64}', key) is not None)
            require(isinstance(aliases, dict) and set(aliases) <= set(columns) and all(isinstance(v, str) and v for v in aliases.values()))
            seen, previous, count = set(), 1, 0
            letters = [get_column_letter(i + 1) for i in range(len(columns))]
            prefix = f'{filename} / {sheet}!'
            for line in stream:
                item = json.loads(line)
                require(isinstance(item, dict))
                if 'end' in item:
                    require(set(item) == {'end'} and type(item['end']) is int and item['end'] == count and count > 0)
                    require(not stream.read())  # gzip CRC와 잘린 파일도 끝까지 확인합니다.
                    return
                require(set(item) == {'row', 'values', 'typed'})
                number, values, typed = item['row'], item['values'], item['typed']
                require(type(number) is int and number > previous)
                require(isinstance(values, list) and len(values) == len(columns) and all(isinstance(v, str) for v in values))
                require(isinstance(typed, dict) and set(typed) <= set(columns) and (allow_typed or not typed))
                row = dict(zip(columns, values))
                for column, cell in typed.items():
                    require(isinstance(cell, dict) and set(cell) == {'kind', 'value', 'number_format'})
                    kind, value = cell['kind'], cell['value']
                    require(kind in ('int', 'float', 'bool', 'datetime', 'date', 'time', 'timedelta') and isinstance(cell['number_format'], str))
                    require((kind == 'int' and type(value) is int) or (kind == 'float' and type(value) in (int, float) and math.isfinite(value))
                            or (kind == 'bool' and type(value) is bool) or (kind in ('datetime', 'date', 'time', 'timedelta') and isinstance(value, str)))
                    require(str(value) == row[column])
                row.update(_id_column=id_column, _source_sheet=sheet, _source_key=key, _typed_cells=typed,
                           _excel_location=f'{prefix}A{number}:{letters[-1]}{number}',
                           _cell_locations={column: f'{prefix}{letter}{number}' for column, letter in zip(columns, letters)})
                if aliases:
                    row['_source_columns'] = aliases
                identity = record_id(row)
                require(identity not in seen)
                text_for(row)
                seen.add(identity)
                previous, count = number, count + 1
                yield row
            require(False)  # 완료 행이 없는 파일을 일부 성공으로 처리하지 않습니다.
    except (OSError, EOFError, UnicodeError, json.JSONDecodeError) as error:
        raise ValueError('변환 파일을 읽지 못했습니다. Excel에서 다시 변환하세요.') from error


def read_cache(path, sheet_name='CSList', **options):
    return list(iter_cache(path, sheet_name, **options))
