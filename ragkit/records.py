"""Excel 입력, 원본 보존, 검색용 문자열과 원문 버전 계산."""
import hashlib
import json
from datetime import date, time as daytime, timedelta
from pathlib import Path
from types import SimpleNamespace
from openpyxl.utils.cell import get_column_letter
from openpyxl import load_workbook

def cell_string(value):
    # 검색과 API용 문자열입니다. Excel의 자료형과 원래 값은 _typed_cells에 별도로 보존합니다.
    if value is None:
        return ''
    return value.isoformat() if isinstance(value, (date, daytime)) else str(value)


def public_cells(row):
    return {key: value for key, value in row.items() if not key.startswith('_')}


def record_id(row, id_column=None):
    # 내부 식별자는 원본 열에서 읽습니다. Excel에 id라는 가짜 열을 추가하지 않습니다.
    if not isinstance(row, dict):
        raise ValueError('기록은 열 이름과 값으로 구성되어야 합니다.')
    key = id_column or row.get('_id_column', 'id')
    value = row.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'비어 있지 않은 문자열 {key}가 필요합니다.')
    return value


def row_revision(row):
    # 행 번호가 바뀌어도 검토 결과를 유지하되 같은 행의 문맥·값이 바뀌면 재검토합니다.
    values = [public_cells(row), row.get('_typed_cells', {})]
    if row.get('_id_column', 'id') != 'id':
        values.append(row['_id_column'])
    return hashlib.sha256(json.dumps(values, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def read_excel(path, sheet_name='점검기록', *, required_columns=None, allow_typed=True, id_column='id', reader='openpyxl', header_aliases=None, progress=None):
    # 변환 파일도 동일한 원문·자료형·셀 위치를 반환하여 후속 처리를 공유합니다.
    # 첫 행은 열 이름, 두 번째 행부터는 데이터입니다. 열 순서는 바꿔도 됩니다.
    # 숫자·날짜·참/거짓 셀은 값과 자료형 정보를 별도로 보존하며 LLM 수정 대상에서 제외합니다.
    path = Path(path)
    if path.name.lower().endswith('.rag.jsonl.gz'):
        from .source_cache import read_cache
        return read_cache(path, sheet_name, required_columns=required_columns,
                          allow_typed=allow_typed, id_column=id_column)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('이 데모는 .xlsx 파일을 읽습니다.')
    if reader == 'xlwings':
        from .excel_app import read_open_book
        return read_open_book(path, sheet_name, required_columns, allow_typed, id_column, header_aliases, progress=progress)
    if reader != 'openpyxl':
        raise ValueError('읽기 방식은 openpyxl 또는 xlwings여야 합니다.')
    # data_only=False로 수식 자체를 확인합니다. 계산되지 않은 수식을 빈 값으로
    # 오해하지 않도록 수식 셀을 발견하면 아래에서 중단합니다. 파일을 수정하지 않습니다.
    book = load_workbook(path, data_only=False)
    try:
        if sheet_name is None:
            if len(book.sheetnames) != 1:
                raise ValueError(f'시트가 여러 개입니다. 시트 이름을 지정하세요: {book.sheetnames}')
            sheet_name = book.sheetnames[0]
        if sheet_name not in book.sheetnames:
            raise ValueError(f'시트를 찾을 수 없습니다: {sheet_name}')
        sheet = book[sheet_name]
        if sheet.merged_cells.ranges:
            raise ValueError('입력 시트의 병합 셀을 해제하고 각 행에 값을 입력하세요.')
        return rows_from_cells(path, sheet_name, sheet.iter_rows(), required_columns, allow_typed, id_column, header_aliases)
    finally:
        # 성공하거나 오류가 발생해도 파일 핸들을 닫아 Excel 파일을 계속 편집할 수 있게 합니다.
        book.close()


class HeaderMappingError(ValueError):
    def __init__(self, message, *, missing, unknown, observed):
        super().__init__(message)
        self.missing = list(missing)
        self.unknown = list(unknown)
        self.observed = list(observed)


# CSList 내보내기에서 확인된 표기만 내부 열 이름으로 맞춥니다. 다른 표기는 바꾸지 않습니다.
CSLIST_HEADER_ALIASES = {
    'Part Desc(1)': 'Part Decs(1)',
    'Part Desc(2)': 'Part Decs(2)',
    'UNIT(CIS)': 'UNIT(CS)',
    'Line': 'line',
    '공수1(1)': '공수(1)',
    '공수1(2)': '공수(2)',
    '도급여부 CASE': '도급여부CASE',
    '무상자재개수': '무상자개개수',
    '불출위치(1)': '분출위치(1)',
    '불출위치(2)': '분출위치(2)',
    '요청 납기일': '요청납기일',
    'CS조치결재상태(조치1/조치2)': '조치결재상태(조치1/조치2)',
    '협력시통합정산CS접수번호': '협력사통합전산CS접수번호',
    '협력사통합정산CS접수번호': '협력사통합전산CS접수번호',
}

MAX_HEADER_ALIASES = 400
MAX_HEADER_NAME = 200
ALIASES_FORMAT = 'column-aliases-v1'


def validate_header_aliases(aliases, required_columns, id_column):
    if not isinstance(aliases, dict):
        raise ValueError('열 이름 대응은 객체여야 합니다.')
    if required_columns is None:
        raise ValueError('양식 열이 있어야 열 이름을 대응할 수 있습니다.')
    required_set = set(required_columns)
    if id_column is not None and required_set and id_column not in required_set:
        raise ValueError('식별 열이 양식에 없습니다.')
    if len(aliases) > MAX_HEADER_ALIASES:
        raise ValueError('열 이름 대응이 너무 많습니다.')
    cleaned = {}
    for source, target in aliases.items():
        if not isinstance(source, str) or not source.strip() or source.strip().startswith('_'):
            raise ValueError('대응할 원본 열 이름이 올바르지 않습니다.')
        source = source.strip()
        if len(source) > MAX_HEADER_NAME:
            raise ValueError('대응할 원본 열 이름이 너무 깁니다.')
        if source in required_set:
            raise ValueError(f'양식 열 이름은 다른 열로 바꾸지 않습니다: {source}')
        if source in CSLIST_HEADER_ALIASES:
            raise ValueError(f'이미 정해진 열 이름 대응은 바꾸지 않습니다: {source}')
        if source in cleaned:
            raise ValueError(f'공백을 제외하면 중복된 원본 열입니다: {source}')
        if not isinstance(target, str) or not target.strip() or target.strip().startswith('_'):
            raise ValueError('맞출 양식 열 이름이 올바르지 않습니다.')
        target = target.strip()
        if len(target) > MAX_HEADER_NAME:
            raise ValueError('맞출 양식 열 이름이 너무 깁니다.')
        if target not in required_set:
            raise ValueError(f'맞출 양식 열이 없습니다: {target}')
        cleaned[source] = target
    return cleaned


def load_header_aliases(path, required_columns, id_column, schema_name=None):
    path = Path(path)
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding='utf-8-sig'))
    except OSError as error:
        raise ValueError('열 이름 대응 파일을 읽지 못했습니다.') from error
    except ValueError as error:
        raise ValueError('열 이름 대응 파일이 올바른 JSON이 아닙니다.') from error
    if not isinstance(data, dict) or data.get('format') != ALIASES_FORMAT:
        raise ValueError('열 이름 대응 파일 형식을 확인하세요.')
    if data.get('schema_name') != (schema_name or 'CSList'):
        raise ValueError('저장된 열 이름 대응이 현재 양식과 다릅니다.')
    if data.get('id_column') != id_column or data.get('columns') != list(required_columns):
        raise ValueError('저장된 열 이름 대응이 현재 양식과 다릅니다.')
    if 'aliases' not in data:
        raise ValueError('열 이름 대응 파일 형식을 확인하세요.')
    return validate_header_aliases(data['aliases'], required_columns, id_column)


def save_header_aliases(path, aliases, required_columns, id_column, schema_name=None):
    validated = validate_header_aliases(aliases, required_columns, id_column)
    path = Path(path)
    payload = {
        'format': ALIASES_FORMAT,
        'schema_name': schema_name or 'CSList',
        'id_column': id_column,
        'columns': list(required_columns),
        'aliases': validated,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + '.tmp')
    try:
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        tmp.replace(path)
    except Exception:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise
    return validated


def _map_cslist_headers(headers, required_columns, id_column, header_aliases=None):
    if id_column != 'CS 접수번호' or required_columns is None:
        return headers, None
    required = set(required_columns)
    custom = header_aliases or {}
    mapped, source_columns = [], {}
    for key in headers:
        if key in required:
            mapped.append(key)
            continue
        canonical = CSLIST_HEADER_ALIASES.get(key)
        if canonical is None and key in custom:
            canonical = custom.get(key)
        if canonical is not None and canonical in required:
            mapped.append(canonical)
            source_columns[canonical] = key
        else:
            mapped.append(key)
    if len(set(mapped)) != len(mapped):
        raise ValueError(f'첫 행에는 중복·빈 이름 없이 {id_column}와 내용 열을 입력하세요. _로 시작하는 열은 예약되어 있습니다.')
    return mapped, source_columns or None


def rows_from_cells(path, sheet_name, cells, required_columns, allow_typed, id_column, header_aliases=None):
    return list(iter_rows_from_cells(path, sheet_name, cells, required_columns, allow_typed, id_column, header_aliases))


def iter_rows_from_cells(path, sheet_name, cells, required_columns, allow_typed, id_column, header_aliases=None):
    # 파일 직접 읽기와 Excel 연결 읽기에 같은 검사를 적용합니다. 열·ID·빈칸 규칙이 달라지지 않습니다.
    cells = iter(cells)
    header_cells = next(cells, ())
    headers = [cell.value.strip() if isinstance(cell.value, str) else cell.value
               for cell in header_cells]
    # id 외의 열 이름은 자유롭게 사용할 수 있습니다. LLM은 원래 열 이름을 유지하며 셀별 문장만 다듬습니다.
    if (not headers or any(not isinstance(key, str) or not key or key.startswith('_') for key in headers)
            or len(set(headers)) != len(headers)):
        raise ValueError(f'첫 행에는 중복·빈 이름 없이 {id_column}와 내용 열을 입력하세요. _로 시작하는 열은 예약되어 있습니다.')
    if required_columns is None and headers.count(id_column) != 1:
        raise ValueError(f'첫 행에는 중복·빈 이름 없이 {id_column}와 내용 열을 입력하세요. _로 시작하는 열은 예약되어 있습니다.')
    observed = list(headers)
    aliases = None
    if header_aliases is not None:
        aliases = validate_header_aliases(header_aliases, required_columns, id_column)
    headers, source_columns = _map_cslist_headers(headers, required_columns, id_column, aliases)
    if required_columns is not None and set(headers) != set(required_columns):
        missing = sorted(set(required_columns) - set(headers))
        extra = sorted(set(headers) - set(required_columns))
        raise HeaderMappingError(
            f'첫 행의 열이 양식과 다릅니다. 누락: {missing}, 추가: {extra}',
            missing=missing, unknown=extra, observed=observed)
    if headers.count(id_column) != 1:
        raise ValueError(f'첫 행에는 중복·빈 이름 없이 {id_column}와 내용 열을 입력하세요. _로 시작하는 열은 예약되어 있습니다.')
    seen = set()
    source_key = hashlib.sha256((str(path.resolve()) + '!' + sheet_name).encode()).hexdigest()
    for row_number, cell_row in enumerate(cells, start=2):
        # 읽기 전용 Excel의 빈 셀에도 원래 좌표를 붙입니다. 마지막 빈 열도 보존합니다.
        cell_row = list(cell_row)
        if len(cell_row) > len(headers):
            raise ValueError(f'{sheet_name} {row_number}행: 첫 행에 이름이 없는 열이 있습니다.')
        for index in range(len(headers)):
            if index >= len(cell_row):
                cell_row.append(SimpleNamespace(value=None, data_type='n', number_format='General'))
            if not hasattr(cell_row[index], 'coordinate'):
                cell = cell_row[index]
                cell_row[index] = SimpleNamespace(value=cell.value, data_type=cell.data_type,
                    number_format=cell.number_format, coordinate=f'{get_column_letter(index+1)}{row_number}')
        values = [cell.value for cell in cell_row]
        # 완전히 빈 행은 건너뛰되, 다른 값이 있는데 ID만 빠진 행은 오류로 처리합니다.
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
            continue
        for cell in cell_row:
            if cell.data_type in ('f', 'e'):
                raise ValueError(f'{sheet_name}!{cell.coordinate}: 수식·오류 대신 확정된 문자 값을 입력하세요.')

        if not allow_typed and any(v is not None and not isinstance(v, str) for v in values):
            raise ValueError('이 로컬 데모에서는 문자 값을 입력하세요.')
        # 빈 Excel 셀(None)은 빈 문자열로 바꿉니다. 실제 본문의 공백은 보존합니다.
        row = dict(zip(headers, (cell_string(value) for value in values)))
        row['_id_column'] = id_column
        row['_source_sheet'] = sheet_name
        if source_columns:
            row['_source_columns'] = source_columns
        row['_typed_cells'] = {key: {'kind': type(cell.value).__name__,
                              'value': cell_string(cell.value) if isinstance(cell.value, (date, daytime, timedelta)) else cell.value,
                              'number_format': cell.number_format}
                              for key, cell in zip(headers, cell_row) if cell.value is not None and not isinstance(cell.value, str)}
        row['_source_key'] = source_key
        text_for(row)
        identity = record_id(row)
        if identity in seen:
            raise ValueError(f'{sheet_name} {row_number}행: 중복 ID ({id_column}) {identity}')
        seen.add(identity)
        # 사용자 작성 '출처'는 그대로 두고, 실제 파일·시트·행 위치를 별도로 기록합니다.
        row['_excel_location'] = f'{path.name} / {sheet_name}!{cell_row[0].coordinate}:{cell_row[-1].coordinate}'
        row['_cell_locations'] = {key: f'{path.name} / {sheet_name}!{cell.coordinate}'
                                  for key, cell in zip(headers, cell_row)}
        yield row
    if not seen:
        raise ValueError('Excel에 읽을 데이터 행이 없습니다.')


def text_for(row, id_column=None):
    # 원본 열 이름과 값을 그대로 보존합니다. 작성 내용의 정리는 normalize()에서 수행합니다.
    identity = record_id(row, id_column)
    id_column = id_column or row.get('_id_column', 'id')
    content = [(key, value) for key, value in row.items() if key != id_column and not key.startswith('_')]
    if not content or any(not isinstance(value, str) for _, value in content):
        raise ValueError('내용 열은 문자열이어야 합니다.')
    if not any(value.strip() for _, value in content):
        raise ValueError(f'{identity}: 내용이 비어 있습니다.')
    # CS의 빈 항목은 원본 JSON에 남깁니다. 검색 문서에 빈 열 이름 128개를 반복하지 않습니다.
    text = f"기록 ID: {identity}\n" + '\n'.join(f'{key}: {value}' for key, value in content if id_column == 'id' or value.strip())
    if row.get('_excel_location'):
        text += f"\nExcel 위치: {row['_excel_location']}"
    return text
